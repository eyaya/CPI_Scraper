"""Angola — INE Inquérito sobre o Emprego em Angola (IEA), time-series workbooks.

INE replaced the quarterly IEA bulletin PDF with workbooks split by
METHODOLOGY, and this indicator reads two of them:

* PRIMARY -- "2.QUADROS COMPLEMETARES_IEA (ANTIGA METODOLOGIA_13ª CIET).xlsx",
  2019 Q2 to 2025 Q3 on the 13th-ICLS basis:
    - sheet "Taxa de emprego", block "Sector de actividade económica": shares
      of the employed across INE's ten national branch groups -> `industry`;
    - same sheet, block "Situação perante o emprego": twelve situations that
      mix employer type (administração pública, empresa pública, sector
      privado...) with status (conta própria, trabalho familiar...) -- INE's
      own single variable, so `employment_status`, classification National;
    - sheet "Planilha3" (Quadro 9, III trimestre 2025): employed COUNTS in the
      sector informal / formal, nationally and by urban/rural -> `formality`,
      and the Total column by branch -> `industry` counts.
* EXTRA -- "4.QUADROS COMPLEMETARES_IEA (NOVA METODOLOGIA_19ª, 20ª e 21ª
  CIET).xlsx", 2025 Q4 onwards on the 19th-21st ICLS basis:
    - "Taxa de emprego", "Sector de actividade económica": 21 branches that read
      like ISIC Rev.4 sections, though no scheme is named -> National;
    - "Taxa de emprego", "Situação perante o emprego": six EMPLOYER TYPES
      (governo, fazenda, empresa privada, famílias, ONG, organização
      internacional). INE heads it as the old block, but its categories are
      institutional units, not statuses, so it is filed as `sector`.

THE TWO METHODOLOGIES MUST NEVER BE CHAINED. The 13th-ICLS "employed" counted
own-use producers and ran to 14,3 million in 2025 Q3; the 19th-ICLS basis gives
8,9 million in 2025 Q4. The periods do not overlap, but the `survey` column
names the basis on every row so that no user joins across it unawares.

WHERE THE FIGURES SIT. The sheet is titled "Taxa de emprego" and its first
blocks ARE employment rates (by residence, sex, age), which belong to
`unemployment` and are not read here. The sector and situation blocks below
them are distributions -- every column sums to 100 -- so they are composition.

TWO OLD COLUMNS DO NOT SUM TO 100 and are collected anyway: 2024 I trim
(104,57) and 2024 Anual (101,07), identically in both blocks -- a uniform
inflation consistent with a denominator slip, not a mislabelled row. Their sums
are pinned in `_OFF_SUM`; every other column must sum to 100 within 0,05.

ONE PUBLISHED COLUMN IS NOT COLLECTED: the 2020 "Anual" column of the old
sector block. It sums to 100 but its values sit against the wrong labels --
Agricultura 1,33 and Administração pública 55,75, where every 2020 quarter
has agriculture at 55-59 and administration at 6-7. Collecting it would publish
real numbers under the wrong categories. It is excluded by name and CHECKED to
still be scrambled on every run, so if INE corrects it the parser raises and
the exclusion can be lifted. The status block's 2020 Anual column agrees with
its quarters and is collected.

OBSERVED, NOT RESOLVED: in the new workbook, "Família(s) como trabalhador(a)
doméstico(a)" (situação) equals "Não classificadas noutras categorias"
(sector) to every digit in 2026 Q1 and Q2 but not in 2025 Q4. Both blocks
still sum to 100, so neither can be shown to be the copy; both are kept.

"Planilha3" appears in BOTH workbooks, identically. It is 13th-ICLS: its total
of 14 332 681 employed is exactly the old series' 2025 Q3 figure. It is read
from the primary file only.

BASE is 15+ (Quadro 9's title; the IEA's standard). Shares are INE's unrounded
cell values, kept as stored.

CROSS-CHECK: old 2025 Q3 industry shares — Agricultura 41,25; Comércio 22,25;
Indústria, Energia e água 7,33. Quadro 9: sector informal 11 053 485, formal
3 279 197, total 14 332 681; rural informal 3 898 364. New 2026 Q2 —
Agricultura, silvicultura e pesca 25,96; Comércio 30,00; Empresa privada (não
agrícola) 63,63.
"""
from __future__ import annotations

import os
import re

import openpyxl
import pandas as pd

from ._common import row

BASE = "15+"
SURVEY_OLD = "IEA (antiga metodologia, 13ª CIET)"
SURVEY_NEW = "IEA (nova metodologia, 19ª, 20ª e 21ª CIET)"

_Q = {"i": 1, "ii": 2, "iii": 3, "iv": 4}

# (year, column label) -> block whose column is known to be scrambled.
_SCRAMBLED = {("2020", "Anual"): "Sector de actividade"}

# COLUMNS THAT DO NOT SUM TO 100 AS PUBLISHED, in BOTH old blocks alike:
# 2024 I trim sums to 104,57 and the 2024 Anual column to 101,07 -- about the
# average of that quarter and three that do sum to 100. Every category in the
# quarter reads high by the same few per cent, which points at INE dividing by
# a smaller denominator that quarter, not at a mislabelled row. Collected as
# printed; the published sum is pinned here so any other drift still raises.
_OFF_SUM = {("2024", "I trim"): 104.57, ("2024", "Anual"): 101.07}


def _period(year, label) -> tuple[str, str]:
    lab = str(label).strip()
    m = re.fullmatch(r"(I{1,3}|IV)\s*trim", lab, re.I)
    if m:
        return f"{int(year)}-Q{_Q[m.group(1).lower()]}", "quarterly"
    if lab.lower().startswith("anual"):
        return str(int(year)), "annual"
    raise ValueError(f"unreadable period header {year!r} {label!r}")


def _blocks(ws) -> tuple[list, dict[str, list]]:
    """Header columns [(col, year, label)] and {block heading: [row, ...]}."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    cols = [(i, rows[1][i], rows[2][i]) for i in range(1, len(rows[1]))
            if rows[1][i] is not None and rows[2][i] is not None]
    blocks, cur = {}, None
    for r in rows[3:]:
        lab = r[0]
        if lab is None or str(lab).startswith("Fonte"):
            continue
        if all(r[i] is None for i, _, _ in cols):
            cur = str(lab).strip()
            blocks[cur] = []
        elif cur is not None:
            blocks[cur].append(r)
    return cols, blocks


def _block(blocks: dict, prefix: str) -> list:
    hits = [k for k in blocks if k.lower().startswith(prefix.lower())]
    if len(hits) != 1:
        raise ValueError(f"expected one block {prefix!r}, found {hits}")
    return blocks[hits[0]]


def _shares(path, specs, survey, code) -> list[dict]:
    ws = openpyxl.load_workbook(path, data_only=True)["Taxa de emprego"]
    cols, blocks = _blocks(ws)
    out = []
    for prefix, topic, cls in specs:
        body = _block(blocks, prefix)
        for i, year, label in cols:
            vals = [r[i] for r in body]
            if any(v is None for v in vals):
                continue
            total = sum(vals)
            expect = _OFF_SUM.get((str(year), str(label).strip()), 100)                 if survey == SURVEY_OLD else 100
            if abs(total - expect) > 0.05:
                raise ValueError(f"{survey} {prefix} {year} {label}: sums to "
                                 f"{total:.2f}, not 100")
            period, freq = _period(year, label)
            if _SCRAMBLED.get((str(year), str(label).strip())) == prefix                     and survey == SURVEY_OLD:
                agri = vals[0]
                if agri > 20:
                    raise ValueError(
                        f"{prefix} {year} {label} is no longer scrambled "
                        f"(Agricultura = {agri:.1f}) -- remove it from "
                        f"_SCRAMBLED and collect it")
                continue
            for r, v in zip(body, vals):
                out.append(row(topic=topic, characteristic=str(r[0]).strip(),
                               classification=cls, survey=survey,
                               period=period, reference_period=f"{label} {year}",
                               frequency=freq, working_age_base=BASE,
                               value=v, measure="share", unit="percent",
                               series_code=f"{code} {prefix[:20]}"))
    return out


def _quadro9(path) -> list[dict]:
    ws = openpyxl.load_workbook(path, data_only=True)["Planilha3"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    src = next(str(c) for r in rows for c in r
               if isinstance(c, str) and c.startswith("Fonte"))
    m = re.search(r"\b(I{1,3}|IV)\s+trimestre\s+de\s+(20\d\d)", src, re.I)
    period = f"{m.group(2)}-Q{_Q[m.group(1).lower()]}"
    ref = f"{m.group(1)} trimestre {m.group(2)}"
    common = dict(survey=SURVEY_OLD, period=period, reference_period=ref,
                  frequency="quarterly", working_age_base=BASE,
                  measure="count", unit="persons", series_code="IEA Quadro 9")
    out, block, industries = [], None, []
    for r in rows:
        nums = [v for v in r[4:7] if isinstance(v, (int, float))]
        if len(nums) != 3:
            continue
        inf, frm, tot = nums
        if abs(inf + frm - tot) > 1:
            raise ValueError(f"Quadro 9 {r[1:3]}: {inf} + {frm} != {tot}")
        if r[1]:
            block = str(r[1])
        if r[1] == "Angola":
            loc, loc_lab, cat = "all", "Total", None
        elif "Resid" in (block or ""):
            loc_lab = str(r[2]).strip()
            loc, cat = ("urban" if loc_lab.startswith("Urb") else "rural"), None
        elif "Actividade" in (block or ""):
            loc, loc_lab, cat = "all", "Total", str(r[2]).replace("\xa0", " ").strip()
        else:
            raise ValueError(f"Quadro 9: unknown block {block!r}")
        if cat is None:
            for lab, v in (("Sector informal", inf), ("Sector formal", frm),
                           ("Total", tot)):
                out.append(row(topic="formality", characteristic=lab,
                               classification="Not applicable",
                               locality=loc, locality_label=loc_lab,
                               value=v, **common))
        else:
            industries.append(tot)
            out.append(row(topic="industry", characteristic=cat,
                           classification="National", value=tot, **common))
    national = next(x["value"] for x in out if x["topic"] == "formality"
                    and x["locality"] == "all" and x["characteristic"] == "Total")
    if len(industries) != 10 or abs(sum(industries) - national) > 10:
        raise ValueError(f"Quadro 9: {len(industries)} branches summing to "
                         f"{sum(industries)} against {national}")
    return out


def _pick(paths: list[str], *terms: str) -> str:
    hit = [p for p in paths if all(t.lower() in os.path.basename(p).lower()
                                   for t in terms)]
    if len(hit) != 1:
        raise FileNotFoundError(f"need one workbook matching {terms}: {paths}")
    return hit[0]


def parse(local_path: str, extras: list[str] | None = None) -> pd.DataFrame:
    paths = [local_path] + list(extras or [])
    old = _pick(paths, "QUADROS", "ANTIGA")
    rows = _shares(old, [("Sector de actividade", "industry", "National"),
                         ("Situação perante", "employment_status", "National")],
                   SURVEY_OLD, "IEA13")
    rows += _quadro9(old)
    new = [p for p in paths if "NOVA" in os.path.basename(p).upper()]
    if new:
        rows += _shares(_pick(paths, "QUADROS", "NOVA"),
                        [("Sector de actividade", "industry", "National"),
                         ("Situação perante", "sector", "Not applicable")],
                        SURVEY_NEW, "IEA19")
    return pd.DataFrame(rows)
