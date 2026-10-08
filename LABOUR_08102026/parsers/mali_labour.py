"""Mali — INSTAT EMOP (Enquête Modulaire et Permanente auprès des Ménages),
companion table workbooks ("tab-emop<YY>pas<N>_eq.xlsx"), chapter 4 "Emploi".

NOT THE BULLETIN. `unemployment/` reads INSTAT's labour-market bulletin, which
draws status as an infographic and industry as a three-sector chart -- which is
why Mali was swept as empty here. The EMOP workbooks behind the analysis
reports print the same cuts as TABLES, by region, milieu, sex and education.

A HOUSEHOLD survey: EMOP is INSTAT's continuous household survey, run in
passages; the employment module is fielded in April-June -- passage 2 from 2021,
passage 1 in 2020 -- so each year gives one reading, dated to Q2.

Two tables per workbook:

* Tab4.5  secteur d'activité (Primaire / Industrie / Commerce / Service) -> industry
* Tab4.6  statut salarial (Salariés / Patron, travailleur indépendant /
          Apprenti, aide familiale) -> employment_status

Each ROW is a population group and its values are that group's distribution
across the categories (they sum to 100), so every row is a column-%
distribution in this schema's sense, not a row percentage. The group sits in
the matching column: Région -> geography, Milieu -> locality (the indented
'Bamako / 'Autres Villes rows are urban sub-strata, kept as locality_label),
Sexe -> sex, Niveau d'instruction -> education, Ensemble -> the national row.

PUBLISHED DEFECTS, HANDLED EXPLICITLY, NOT SMOOTHED:

* 2023 TAB4.5'S "ENSEMBLE" ROW IS A COPY OF THE "RURAL" ROW -- 71,74 / 9,9 /
  12,33 / 6,03, digit for digit -- and cannot be the national figure: it lies
  outside the male-female range (59,6 / 56,26 primaire). It is dropped by name
  and re-checked every run (`_KNOWN_COPIES`), so a correction by INSTAT is
  noticed. The same copied row is what the 2024 workbook's Tab4.9 trend prints
  for 2023 (and its 2023 STATUS column copies Rural too), which is why Tab4.9
  is not read at all: each year is taken from its own workbook.
* A GENERAL GUARD backs that up: an Ensemble value outside the range of its own
  Masculin and Féminin values raises, so a new copy of this kind stops the run
  instead of shipping.
* 2022's Tab4.5 source line reads "EMOP 2021" although its values are 2022's
  (they differ from the 2021 workbook throughout). The year therefore comes
  from INSTAT's FILE NAME (tab-emop22pas2) and is checked against the source
  lines of the workbook as a whole -- a majority, so one slip cannot re-date a
  table and a wrong file name cannot pass silently.

REGION NAMES ARE KEPT AS PRINTED, and INSTAT spells them differently across
years (Taoudénit / Taoudenit / Taoudenni, Ménaka / Menaka, Ségou / Segou). The
region set also changes: 9 regions to 2021, 11 in 2022-23, 20 from 2024 (the
new administrative regions). A user joining regions across years must map
them; the collector does not.

COVERAGE CAVEAT (the workbook's own "Infos" sheet): Kidal, Ménaka, Nara,
Douentza and Bandiagara are represented by only part of their sampled clusters
(under 50% / 60% coverage, for security reasons) and are to be read with
caution.

NO NOMENCLATURE IS NAMED -- the four "secteurs" and three "statuts" are
INSTAT's own groupings -> National. Base 15+ (chapter 4's population, Tableau
4.1 "population de 15 ans et plus").

NOT COLLECTED: Tab4.1 (activity status of the 15+, whose rates `unemployment`
holds from the bulletin), Tab4.4 (children 5-17), Tab4.7 (days worked -- no
topic), Tab4.8/4.9 (indicator digests; 4.9 repeats earlier years and carries
the copied 2023 rows), and the "Total" column -- a formula sum over each
row, stored with float residue (100.00000000000254), not a published figure.

CROSS-CHECK (Ensemble): 2025 primaire 58,65 / salariés 10,69; 2024 primaire
55,10; 2022 primaire 52,11; 2021 primaire 68,26; 2020 primaire 63,3 /
salariés 9,8.
"""
from __future__ import annotations

import os
import re
from collections import Counter

import openpyxl
import pandas as pd

from . import _common as C

_SURVEY = "Enquête Modulaire et Permanente auprès des Ménages (EMOP)"
_TABLES = {  # sheet -> (topic, caption must contain)
    "Tab4.5": ("industry", "secteur d"),
    "Tab4.6": ("employment_status", "statut salarial"),
}
_SECTIONS = {"région": "geography", "milieu": "locality", "sexe": "sex",
             "niveau d'instruction": "education"}
_FILE_RE = re.compile(r"tab-emop(\d\d)pas(\d)", re.I)
_SRC_RE = re.compile(r"EMOP[\s-]*(20\d\d),?\s*passage\s*(\d)\s*\(([^)]*)\)", re.I)

# (year, sheet, row label, row it copies) -- verified against the workbook
# every run; if INSTAT corrects it, the check below raises so the drop is
# removed deliberately rather than silently hiding a now-good row.
_KNOWN_COPIES = {("2023", "Tab4.5", "Ensemble", "Rural")}


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s).replace("’", "'")).strip()


def _label(s) -> str:
    # "     'Bamako" / "‘15 à 24 ans": an indent and a stray quote mark
    return _norm(s).lstrip("'‘").strip()


def _year_and_reference(path: str, wb) -> tuple[str, str]:
    m = _FILE_RE.search(os.path.basename(path))
    if not m:
        raise ValueError(f"EMOP: cannot date {path!r} -- expected INSTAT's "
                         f"tab-emop<YY>pas<N> file name")
    year = f"20{m.group(1)}"
    hits = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and v.strip().startswith("Source"):
                    s = _SRC_RE.search(v)
                    if s:
                        hits.append((s.group(1), s.group(2), s.group(3).strip()))
    if not hits:
        raise ValueError(f"EMOP {year}: no source line in the workbook")
    (y, passage, months), _ = Counter(hits).most_common(1)[0]
    if y != year or passage != m.group(2):
        raise ValueError(f"EMOP: file {os.path.basename(path)} says {year} "
                         f"passage {m.group(2)}, its source lines say {y} "
                         f"passage {passage}")
    if not re.match(r"avril\s*-\s*juin", months, re.I):
        raise ValueError(f"EMOP {year}: employment module fielded {months!r}, "
                         f"not April-June -- re-check the period")
    return year, f"EMOP {year}, passage {passage} ({months.lower()})"


def _read_sheet(ws, year: str, reference: str) -> list[dict]:
    topic, must = _TABLES[ws.title]
    rows = [r for r in ws.iter_rows(values_only=True)
            if any(v is not None for v in r)]
    caption = next(v for v in rows[0] if v is not None)
    if must not in _norm(caption).lower():
        raise ValueError(f"EMOP {year} {ws.title}: unexpected caption "
                         f"{caption!r}")
    # THE TABLE DOES NOT ALWAYS START IN COLUMN A: 2021 puts its labels in B
    # and its caption in A or B by sheet. The label column is wherever the
    # first section heading ("Région") sits; everything left of it is dropped.
    lab_col = next(r.index(v) for r in rows[1:] for v in r
                   if isinstance(v, str) and _norm(v).lower() == "région")
    rows = [rows[0]] + [r[lab_col:] for r in rows[1:]]
    # Header: the category names, after a stub cell ("Région/Milieu de
    # résidence", "Caractéristiques sociodémographiques") that 2020-21 omit.
    # Categories are matched to the value columns of the data rows BY ORDER,
    # and the two counts must agree.
    cats = [_norm(v) for v in rows[1] if v is not None
            and not re.match(r"(Région|Caractéristiques)", _norm(v))]
    first = next(r for r in rows[2:]
                 if any(isinstance(v, (int, float)) for v in r))
    vcols = [i for i, v in enumerate(first) if isinstance(v, (int, float))]
    if len(vcols) != len(cats):
        raise ValueError(f"EMOP {year} {ws.title}: {len(cats)} headers {cats} "
                         f"for {len(vcols)} value columns")
    cols = list(zip(vcols, cats))

    section, parsed = None, []
    for r in rows[2:]:
        lab = r[0]
        if lab is None or _norm(lab).startswith("Source"):
            continue
        vals = [r[i] if i < len(r) else None for i, _ in cols]
        key = _norm(lab).lower()
        if all(v is None for v in vals):
            if key not in _SECTIONS:
                raise ValueError(f"EMOP {year} {ws.title}: unknown section {lab!r}")
            section = _SECTIONS[key]
            continue
        # The national row closes the table: it belongs to no section, though
        # it follows the education block directly.
        if key == "ensemble":
            section = None
        parsed.append((section, _label(lab), vals))

    # Copies: drop the known one after proving it is still a copy.
    by_label = {(s, l): v for s, l, v in parsed}
    for (y, sheet, lab, of) in _KNOWN_COPIES:
        if y == year and sheet == ws.title:
            if by_label.get((None, lab)) != by_label.get(("locality", of)):
                raise ValueError(f"EMOP {year} {sheet}: {lab!r} no longer copies "
                                 f"{of!r} -- INSTAT corrected it; remove it from "
                                 f"_KNOWN_COPIES")
            parsed = [p for p in parsed if not (p[0] is None and p[1] == lab)]

    # An Ensemble outside its own sexes' range is not a national figure.
    ens = next((v for s, l, v in parsed if s is None and l == "Ensemble"), None)
    male = by_label.get(("sex", "Masculin"))
    female = by_label.get(("sex", "Féminin"))
    if ens and male and female:
        for (_, cat), e, a, b in zip(cols, ens, male, female):
            if cat != "Total" and not (min(a, b) - 0.05 <= e <= max(a, b) + 0.05):
                raise ValueError(f"EMOP {year} {ws.title}: Ensemble {cat} {e} lies "
                                 f"outside Masculin {a} / Féminin {b}")

    out = []
    for section, lab, vals in parsed:
        ctx = {}
        if section == "geography":
            ctx["geography"] = lab
        elif section == "locality":
            ctx["locality"] = "rural" if lab == "Rural" else "urban"
            ctx["locality_label"] = lab
        elif section == "sex":
            ctx["sex"] = C.normalise_sex(lab)
        elif section == "education":
            ctx["education"] = lab
        elif lab != "Ensemble":
            raise ValueError(f"EMOP {year} {ws.title}: row {lab!r} outside any section")
        for (_, cat), v in zip(cols, vals):
            # THE "TOTAL" COLUMN IS NOT COLLECTED: it is an Excel SUM over the
            # row (100.00000000000254 -- float residue, not a published
            # figure), printed in Tab4.6 and only 2023's Tab4.5.
            if v is None or cat == "Total":
                continue
            out.append(C.row(topic=topic, characteristic=cat,
                             classification="National", value=v,
                             survey=_SURVEY, period=f"{year}-Q2",
                             reference_period=reference, frequency="annual",
                             measure="share", unit="percent",
                             working_age_base="15+",
                             series_code=f"EMOP {ws.title}", **ctx))
    return out


def _parse_one(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    year, reference = _year_and_reference(path, wb)
    missing = [s for s in _TABLES if s not in wb.sheetnames]
    if missing:
        raise ValueError(f"EMOP {year}: sheets {missing} not in the workbook")
    out = []
    for sheet in _TABLES:
        got = _read_sheet(wb[sheet], year, reference)
        # region / milieu / sexe / instruction / Ensemble: at least 9 + 4 + 2
        # + 5 rows of categories each (fewer = a section went unread).
        if len({(r["geography"], r["locality_label"], r["sex"], r["education"])
                for r in got}) < 20:
            raise ValueError(f"EMOP {year} {sheet}: only {len(got)} values read")
        out += got
    return out


def parse(path: str, extras: list[str] | None = None) -> pd.DataFrame:
    rows = []
    for p in [path, *(extras or [])]:
        rows += _parse_one(p)
    return pd.DataFrame(rows)
