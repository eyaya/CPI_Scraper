"""Morocco — HCP workbook "Principaux indicateurs trimestriels rétropolés
provisoires du marché du travail selon la méthodologie EMO, Période de 2017 à
2025" (https://www.hcp.ma/file/248688/). Added 2026-10-12.

HCP's own back-cast of the 2017-Q1..2025-Q4 quarters onto the NEW Enquête sur
la Main-d'Œuvre (EMO) basis that replaced the ENE in 2026. It is a different
methodology from the ENE annual report that `morocco_unemployment` reads
(2025 ENE chômage 13,0% against EMO-basis 2025 quarters of 10,3-11,8%), so it
is carried as its own survey and its own series labels -- never chained to
the ENE rows, and the ENE 2025 quarters are NOT replaced. HCP marks the
figures "provisoires"; they are kept exactly as published.

Sheet "Données": one row per quarter (AN, TRIM), one column per indicator
code; sheet "Métadonnées" names each code and its unit (% or Millier). Codes
are TX.CH.STR (taux de chômage strict), TX.EMP.REV (taux d'emploi contre
revenu -- employment for pay or profit, the 19th ICLS concept), TX.PART.MO
(taux de participation à la main-d'œuvre), POP.CH.STR / POP.EMP.REV / MO
(levels, thousands), each also by one breakdown: .U / .R (milieu), .H / .F
(sexe), 15-24 / 25-34 / 35+ (âge). No cross-tabulations.

CROSS-CHECK: 2017-Q1 chômage 10,0 (Urbain 13,4, Rural 4,8; 15-24 ans 21,3);
2025-Q4 chômage 11,5, main-d'œuvre 11 764 mille = 1 349 + 10 415.
"""
from __future__ import annotations

import re

import openpyxl

SURVEY = ("Enquête sur la Main-d'Œuvre (EMO) — séries trimestrielles "
          "rétropolées provisoires 2017-2025 (HCP)")
_TAG = " [EMO, rétropolé provisoire]"
_STEM = {
    "TX.CH.STR": ("unemployment_rate", "strict", "rate", "percent"),
    "TX.EMP.REV": ("employment_to_population_ratio", "not_applicable", "rate", "percent"),
    "TX.PART.MO": ("labour_force_participation_rate", "strict", "rate", "percent"),
    "POP.CH.STR": ("unemployed", "strict", "count", "thousand_persons"),
    "POP.EMP.REV": ("employed", "not_applicable", "count", "thousand_persons"),
    "MO": ("labour_force", "not_applicable", "count", "thousand_persons"),
}
_CODE = re.compile(r"^(TX\.CH\.STR|TX\.EMP\.REV|TX\.PART\.MO|POP\.CH\.STR|POP\.EMP\.REV|MO)"
                   r"(?:\.([URHF])|(15-24|25-34|35\+))?$")
_UNIT = {"%": "percent", "Millier": "thousand_persons"}


def parse_retro(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    meta = {r[0]: (r[1], r[2]) for r in wb["Métadonnées"].iter_rows(min_row=2, values_only=True)
            if r and r[0]}
    rows = list(wb["Données"].iter_rows(values_only=True))
    head = list(rows[0])
    if head[:2] != ["AN", "TRIM"] or len(head) != 50:
        raise ValueError(f"HCP retro: header {head[:4]}... ({len(head)} columns)")
    out = []
    for code in head[2:]:
        m = _CODE.match(code)
        if not m or code not in meta:
            raise ValueError(f"HCP retro: unknown column {code!r}")
        topic, definition, measure, unit = _STEM[m.group(1)]
        if _UNIT.get(meta[code][1]) != unit:
            raise ValueError(f"HCP retro: {code} unit {meta[code][1]!r}")
        split, age = m.group(2), m.group(3)
        sex = {"H": "male", "F": "female"}.get(split, "total")
        loc, loc_label = {"U": ("urban", "Urbain"), "R": ("rural", "Rural")}.get(
            split, ("all", "Total"))
        t = "youth_unemployment_rate" if topic == "unemployment_rate" and age == "15-24" else topic
        col = head.index(code)
        for r in rows[1:]:
            if r[0] is None:
                continue
            year, q, v = int(r[0]), int(r[1]), r[col]
            if v is None:
                continue
            out.append({
                "topic": t, "definition": definition,
                "series_label": meta[code][0] + _TAG,
                "sex": sex, "age_group": age or "Total", "education": "Total",
                "geography": "Total country", "locality": loc,
                "locality_label": loc_label, "working_age_base": "15+",
                "period": f"{year}-Q{q}",
                "reference_period": f"{q}{'er' if q == 1 else 'ème'} trimestre {year}",
                "frequency": "quarterly", "measure": measure, "value": float(v),
                "unit": unit, "survey": SURVEY, "series_code": f"HCP EMO-R {code}"})
    if len(out) != 48 * 36:
        raise ValueError(f"HCP retro: {len(out)} cells, expected {48 * 36}")
    return out
