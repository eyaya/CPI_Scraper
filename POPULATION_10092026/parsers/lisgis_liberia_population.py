"""LISGIS Liberia 2022 census — population by county and sex (clean Excel).

Parser for the LISGIS county_population.xlsx (2022 Population & Housing Census):
one row per county with Total / Male / Female. We emit those as published (age not
broken down in this file) plus a national "Total country" row summed from the 15
counties (an exhaustive-parts aggregation of the census counts).
"""
from __future__ import annotations
import json

import pandas as pd
from openpyxl import load_workbook


def _num(v):
    if v is None:
        return None
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def _rows(geo, m, f, t):
    return [{"series_type": "census", "sex": sex, "age_group": "Total",
             "geography": geo, "period": "2022", "frequency": "annual",
             "measure": "count", "value": v, "unit": "persons",
             "series_code": "LR_PHC2022"} for sex, v in
            (("total", t), ("male", m), ("female", f)) if v is not None]


def parse(local_path: str) -> pd.DataFrame:
    wb = load_workbook(local_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    out, sm, sf, st = [], 0.0, 0.0, 0.0
    for r in rows[1:]:
        name = str(r[0]).strip() if r and r[0] is not None else ""
        total, male, female = _num(r[1]), _num(r[2]), _num(r[3])
        if not name or total is None:
            continue
        out += _rows(name, male, female, total)
        st += total
        sm += male or 0
        sf += female or 0
    if not out:
        raise ValueError("lisgis_liberia_population: no county rows parsed")
    # national total = sum of the (exhaustive) census counties
    out += _rows("Total country", sm, sf, st)
    return pd.DataFrame.from_records(out)


_AGE_SETS = {"census-population-age-sex-residence",
             "census-population-age-sex-residence-by-county"}


def _age_rows(path: str, names: dict) -> list[dict]:
    """Five-year age x sex from the census THEMATIC REPORT on population size,
    distribution and structure, served by the dataset API (added 2026-10-06):
    the national table and Tables 47-61 (one per county).

    ONLY THE AGE BANDS ARE EMITTED. Each table's own "Total" row would share a
    merge key with the county table's all-ages row, and the two LISGIS tables
    do not agree everywhere: they print the same county totals but different
    sex splits for Grand Cape Mount (county table 96,746 M / 82,121 F; thematic
    96,757 / 82,110) and Lofa (183,063 / 184,238 -- which sum to 367,301, not
    the 367,376 the same row prints -- against 183,100 / 184,276, which do).
    The all-ages series therefore stays the county table's, unchanged, and
    every band set is held to its own table's Total instead.

    Urban/rural columns are not collected: this indicator has no locality
    column, the same deferral as Namibia's.
    """
    with open(path, encoding="utf-8") as fh:
        doc = json.load(fh)
    if doc.get("id") not in _AGE_SETS or str(doc.get("version")) != "2022":
        raise ValueError(f"{path}: not a 2022 LISGIS age-sex dataset")
    by_geo: dict[str, dict[str, dict]] = {}
    for r in doc["rows"]:
        geo = "Total country" if "geo_code" not in r else names[r["geo_code"]]
        by_geo.setdefault(geo, {})[r["age_group"]] = r
    out = []
    for geo, ages in by_geo.items():
        total = ages.pop("Total", None)
        if total is None or len(ages) != 17:
            raise ValueError(f"LISGIS age table {geo}: expected 17 bands + Total")
        for sex, col in (("total", "total_population"), ("male", "total_male"),
                         ("female", "total_female")):
            if sum(_num(r[col]) or 0 for r in ages.values()) != _num(total[col]):
                raise ValueError(f"LISGIS age table {geo} {sex}: bands do not "
                                 f"sum to the table's Total")
            for band, r in ages.items():
                v = _num(r[col])
                if v is None:
                    continue
                out.append({"series_type": "census", "sex": sex,
                            "age_group": band, "geography": geo,
                            "period": "2022", "frequency": "annual",
                            "measure": "count", "value": v, "unit": "persons",
                            "series_code": "LR_PHC2022_THEMATIC"})
    return out


def parse_api(local_path: str, extras: list[str] | None = None) -> pd.DataFrame:
    """The same county table, from LISGIS's dataset API.

    LISGIS rebuilt its site in 2026: county_population.xlsx is gone and the old
    path answers 200 with the app's HTML shell. /api/datasets/population-counties
    serves the 2022 census county table as JSON -- a row per county code with
    population / male / female, and the code-to-name list in
    `dimension_attributes`. Output is identical to the workbook's (Bomi 133,705
    = 68,574 + 65,131), so the series continues unbroken.
    """
    with open(local_path, encoding="utf-8") as fh:
        doc = json.load(fh)
    if doc.get("id") != "population-counties":
        raise ValueError(f"{local_path}: not the population-counties dataset")
    names = {code: a["county_name"] for code, a in
             doc["dimension_attributes"]["county_code"].items()}
    out, sm, sf, st = [], 0.0, 0.0, 0.0
    for r in doc["rows"]:
        if int(r["year"]) != 2022:
            raise ValueError(f"unexpected census year {r['year']}")
        name = names[r["county_code"]]
        total, male, female = (_num(r["population"]), _num(r["male"]),
                               _num(r["female"]))
        if total is None:
            raise ValueError(f"{name}: no population value")
        out += _rows(name, male, female, total)
        st += total
        sm += male or 0
        sf += female or 0
    if len(doc["rows"]) != 15:
        raise ValueError(f"expected Liberia's 15 counties, got {len(doc['rows'])}")
    out += _rows("Total country", sm, sf, st)
    for p in extras or []:
        out += _age_rows(p, names)
    return pd.DataFrame.from_records(out)
