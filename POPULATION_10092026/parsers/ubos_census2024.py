"""Uganda Bureau of Statistics -- National Population and Housing Census 2024,
"Subcounty Profiles" Excel tables, Table 1: population by sex and geographical
location.

Added 2026-10-06 as an extra of the UBOS projection source. Table 1 nests four
levels in ONE column -- district, county, sub-county, parish -- told apart only
by the cell's INDENT (0, 1, 2, 3): the names alone cannot do it, since a
district and one of its sub-counties often share a name (ABIM / ABIM). Only
the indent-0 rows are read:

* the 147 district-level units, names as printed (upper case), and
* the "National" row, as "Total country" (45,905,417).

Sub-county and parish rows are not collected: their names collide with the
districts' and with each other across districts, so as `geography` strings they
would share merge keys with different places.

APAA (9,456) is printed at district level although it is not one of the
districts; it is collected as published. The district-level rows sum EXACTLY
to the National row (checked every run), so APAA is part of the published
national count, not an addition to it. Every row must also satisfy
Male + Female = Total.

series_type = census, period 2024.
"""
from __future__ import annotations

from openpyxl import load_workbook

_SHEET = "Table1"
_TITLE = "Population distribution by Sex and Geographical Location"


def is_census_workbook(path: str) -> bool:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return False
    return _SHEET in wb.sheetnames and "Table_of_Contents" in wb.sheetnames


def parse_census(path: str) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[_SHEET]
    head = " ".join(str(c.value) for row in ws.iter_rows(max_row=3) for c in row
                    if c.value)
    if _TITLE.lower() not in head.lower():
        raise ValueError(f"UBOS NPHC 2024 Table1: unexpected title {head[:90]!r}")
    units, national = [], None
    for row in ws.iter_rows(min_row=4):
        cell = row[0]
        name = str(cell.value).strip() if cell.value is not None else ""
        if not name or not isinstance(row[3].value, (int, float)):
            continue
        if (cell.alignment.indent if cell.alignment else None) != 0:
            continue
        m, f, t = (float(row[i].value) for i in (1, 2, 3))
        if m + f != t:
            raise ValueError(f"UBOS NPHC 2024 {name}: {m} + {f} != {t}")
        if name.lower() == "national":
            national = (m, f, t)
        else:
            units.append((name, m, f, t))
    if national is None or len(units) < 140:
        raise ValueError(f"UBOS NPHC 2024 Table1: national row "
                         f"{'missing' if national is None else 'found'}, "
                         f"{len(units)} district-level units")
    for i, label in enumerate(("male", "female", "total")):
        s = sum(u[i + 1] for u in units)
        if s != national[i]:
            raise ValueError(f"UBOS NPHC 2024: district {label} sum {s} != "
                             f"National {national[i]}")
    out = []
    for name, m, f, t in units + [("Total country", *national)]:
        for sex, v in (("male", m), ("female", f), ("total", t)):
            out.append({"series_type": "census", "sex": sex, "age_group": "Total",
                        "geography": name, "period": "2024",
                        "frequency": "annual", "measure": "count", "value": v,
                        "unit": "persons", "series_code": "UBOS_NPHC2024_T1"})
    return out
