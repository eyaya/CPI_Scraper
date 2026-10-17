"""INE Angola — population projections 2014-2050 and the 1940-2013 series.

Three workbooks from INE's "População e Sociedade / População / Projecções da
População" folder (files are POSTed, like every INE Angola download):

* POPULAÇÃO PROJECTADA POR PROVÍNCIAS, SEGUNDO O SEXO 2014-2050 (Quadro 1) —
  Angola and its 18 provinces, every year 2014-2050, each year a 9-column block
  (Total / Urbana / Rural, each Total / Homens / Mulheres);
* POPULAÇÃO PROJECTADA POR GRUPOS ETÁRIOS 2014-2050 — one sheet per province
  plus "Angola", the same 9-column year blocks down INE's age groups (5-year
  bands, the sub-bands "Menos de 1 ano"/"1-4 anos", broad groups "0-14 anos",
  "15-64 anos", "18 ou mais anos"... and "Idade Mediana");
* POPULAÇÃO 1940_2025 — national population by year.

WHAT IS EMITTED. Only the TOTAL block of each year: this schema has no
residence dimension, so the urban/rural split is dropped rather than encoded
into `geography`. Age bands are INE's own labels, overlapping groups included
(they are separate published series, never summed here). "Idade Mediana" is a
`median_age` in years. The 18 provinces are the PRE-2024 administrative map the
projection was built on; the 2024 census (sources/angola.yaml) uses 21 — the
two are different series and are not reconciled.

The province totals appear in both workbooks; they are taken from Quadro 1 and
the age sheets' total rows are CHECKED equal to it, not emitted twice.

THE 1940-2025 SERIES is not labelled by origin. From 2014 it repeats the
projection's national totals (2015: 26 681 590 in both), so only 1940-2013 is
taken. 1940, 1950, 1960 and 1970 are Angola's colonial census years and their
values are whole counts; 1980 and 1990-2013 are INE estimates (their values
carry decimals, i.e. model output). They are typed `census` and `estimate`
accordingly, and the reasoning is recorded here so it can be revisited.

series_type for 2014-2050 is `projection` (INE's own word, "projectada").

CROSS-CHECK: Angola 2014 25 901 182 (M 12 553 707 / F 13 347 475); Luanda 2014
6 979 233; 1940 3 738 010; Angola 2014 median age 16.
"""
from __future__ import annotations
import re
import openpyxl
import pandas as pd

_SC_PROV = "AO_INE_PROJ_PROV"
_SC_AGE = "AO_INE_PROJ_AGE"
_SC_HIST = "AO_INE_POP_1940_2025"
_CENSUS_YEARS = {1940, 1950, 1960, 1970}
_SEXES = (("total", 0), ("male", 1), ("female", 2))


def _geo(label: str) -> str:
    return "Total country" if label.strip() == "Angola" else label.strip()


def _year_blocks(rows) -> dict[int, int]:
    """year -> first column of its TOTAL block (the Total M/F triple).

    The year headers are merged cells and are not reliably aligned: Quadro 1
    puts "Angola 2014" over its block, the age workbook puts it over the LABEL
    column one to the left. So blocks are found from the sub-headers instead --
    every Total/Homens/Mulheres triple whose residence label (row 2, on the
    block's first or second column) is "Total" -- and paired IN ORDER with the
    header years, which must be the 37 consecutive years 2014-2050."""
    header_years = [int(m.group(1)) for v in rows[1]
                    for m in [re.search(r"(\d{4})", str(v or ""))] if m]
    starts = [c for c in range(1, len(rows[3]) - 2)
              if [str(rows[3][c + k]).strip() for k in range(3)]
              == ["Total", "Homens", "Mulheres"]
              and "Total" in (str(rows[2][c]).strip(), str(rows[2][c + 1]).strip())]
    if header_years != list(range(2014, 2051)) or len(starts) != 37:
        raise ValueError(f"INE Angola: {len(header_years)} header years, "
                         f"{len(starts)} Total blocks (want 37 each)")
    return dict(zip(header_years, starts))


def _row(series, sex, age, geo, period, value, sc, measure="count", unit="persons"):
    return {"series_type": series, "sex": sex, "age_group": age, "geography": geo,
            "period": str(period), "frequency": "annual", "measure": measure,
            "value": float(value), "unit": unit, "series_code": sc}


def _provinces(path: str):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = [r for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    years = _year_blocks(rows)
    totals, out = {}, []
    for r in rows[4:]:
        if not r[0] or str(r[0]).startswith("Fonte") or r[1] is None:
            continue
        geo = _geo(str(r[0]))
        for y, c in years.items():
            t, m, f = r[c], r[c + 1], r[c + 2]
            if abs(m + f - t) > 2:
                raise ValueError(f"INE Angola Q1 {geo} {y}: {m}+{f} != {t}")
            totals[(geo, y)] = (t, m, f)
            for sex, k in _SEXES:
                out.append(_row("projection", sex, "Total", geo, y, (t, m, f)[k], _SC_PROV))
    if len({g for g, _ in totals}) != 19:
        raise ValueError("INE Angola Q1: expected Angola + 18 provinces")
    return out, totals


# PUBLISHED DEFECTS in the age workbook, pinned by (sheet, row label, year) and
# skipped -- any other bad cell still raises:
#   * Cuanza Norte "18 ou mais anos" 2027 reads 18 / 17 / 18 (a stray cell, the
#     neighbouring years are in the hundreds of thousands);
#   * Bié "1-4 anos" 2033 and the Cuanza Sul total row 2014 are typed as text
#     ("345 374", "1 887 710") -- the Cuanza Sul total is in Quadro 1 anyway.
_DEFECTS = {("Cuanza Norte", "18 ou mais anos", 2027), ("Bié", "1-4 anos", 2033),
            ("Cuanza Sul", "Cuanza Sul", 2014)}
# Benguela's sheet breaks the block layout from 2044 (an unlabelled block, then
# 2047-2050 with the triple reordered Homens/Mulheres/Total). Rather than guess
# columns, the sheet is skipped: Benguela's totals still come from Quadro 1.
# The "Moxico" and "Cuando Cubango" sheets carry EACH OTHER's figures: for all
# 37 years each sheet's total row equals the other province in Quadro 1 (Moxico
# sheet 2014 = 535 838 = Cuando Cubango in Quadro 1). Relabelling would be our
# guess, so both are skipped; their totals still come from Quadro 1.
_SKIP_SHEETS = {"Benguela", "Moxico", "Cuando Cubango"}


def _num_ok(v):
    return isinstance(v, (int, float))


def _ages(path: str, totals: dict):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for sheet in wb.sheetnames:
        if sheet.strip() in _SKIP_SHEETS:
            continue
        rows = [r for r in wb[sheet].iter_rows(values_only=True)]
        years = _year_blocks(rows)
        for r in rows[4:]:
            lab = str(r[0] or "").strip()
            if not lab or lab.startswith("Fonte") or r[1] is None:
                continue
            for y, c in years.items():
                vals = (r[c], r[c + 1], r[c + 2])
                if (sheet.strip(), lab, y) in _DEFECTS:
                    continue
                if not all(_num_ok(v) for v in vals):
                    raise ValueError(f"INE Angola {sheet} {lab} {y}: non-numeric {vals}")
                if lab == sheet or _geo(lab) == _geo(sheet):     # the total row
                    if totals.get((_geo(sheet), y)) != vals:
                        raise ValueError(f"INE Angola {sheet} {y}: total {vals} "
                                         f"!= Quadro 1 {totals.get((_geo(sheet), y))}")
                    continue
                if lab.lower().startswith("idade mediana"):
                    for sex, k in _SEXES:
                        out.append(_row("projection", sex, "Total", _geo(sheet), y,
                                        vals[k], _SC_AGE, "median_age", "years"))
                    continue
                if abs(vals[1] + vals[2] - vals[0]) > 2:
                    raise ValueError(f"INE Angola {sheet} {lab} {y}: sexes != total")
                for sex, k in _SEXES:
                    out.append(_row("projection", sex, lab, _geo(sheet), y, vals[k], _SC_AGE))
    return out


def _history(path: str):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for r in wb[wb.sheetnames[0]].iter_rows(values_only=True):
        if isinstance(r[0], (int, float)) and isinstance(r[1], (int, float)):
            y = int(r[0])
            if y < 2014:                       # 2014+ repeats the projection
                out.append(_row("census" if y in _CENSUS_YEARS else "estimate",
                                "total", "Total", "Total country", y, r[1], _SC_HIST))
    if len(out) < 25:
        raise ValueError(f"INE Angola 1940-2025: read {len(out)} years")
    return out


def parse(local_path: str, extras: list[str] | None = None) -> pd.DataFrame:
    """Primary = the provinces workbook (Quadro 1); extras = the age workbook
    and the 1940-2025 series, recognised by file name."""
    out, totals = _provinces(local_path)
    for p in extras or []:
        name = p.upper()
        if "GRUPOS" in name:
            out += _ages(p, totals)
        elif "1940" in name:
            out += _history(p)
    return pd.DataFrame(out)
