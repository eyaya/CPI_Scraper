"""ANStat Côte d'Ivoire -- "Chiffres clés de la Côte d'Ivoire" (Juillet 2025),
the data workbook behind it: the DÉMOGRAPHIE block, 2015-2025. A SECOND
Côte d'Ivoire descriptor (cote_divoire_chiffres_cles.yaml) beside
cote_divoire.yaml (the 2014 and 2021 censuses).

Sheet ENSEMBLE, headed "(p)=projection; (e)=estimation; NB: Seules les données
démographiques de 2021 sont des valeurs réelles", block "DÉMOGRAPHIE(p)",
source column "ANStat/RGPH2021" (fertility and life expectancy
"ANStat/EDS2021", "ANStat/EDS/RP"). Read, every year except 2021:
  Population de la Côte d'Ivoire / Population homme / Population femme
  -> count (persons), sex total / male / female;
  Population de 0-14 ans / 15-59 ans / 60 ans et plus -> count by age_group;
  Taux de croissance démographique -> growth_rate (percent);
  Densité de la population -> density (per_km2);
  Rapport de masculinité -> sex_ratio (ratio);
  Indice synthétique de fécondité (ISF) -> fertility_rate (births_per_woman);
  Espérance de vie à la naissance (Total / hommes / femmes) -> life_expectancy.
series_type projection (the block's "(p)").

2021 IS NOT FILED: it is the RGPH 2021 census, collected by cote_divoire.yaml;
the workbook's 2021 total, male and female are CHECKED equal to the census
figures (29 389 150; 15 344 990; 14 044 160).

AS DISPLAYED. The cells hold unrounded model output (25 476 127.186...);
each value is rounded to its cell's number format -- "#,##0" to the unit,
"0.0" to one decimal, "0.00" to two -- which is what the workbook shows and
what the printed "Chiffres clés" PDF prints (2015: 25 476 127).

NOT READ: urban/rural shares (no locality column), household counts and
size, literacy, and the crude birth rate (no schema measure); sheet DSDS of
the same workbook, whose population figures DIFFER from ENSEMBLE's for every
projected year (2020: 28 572 171 vs 28 725 589; 2022: 30 153 010 vs
30 211 973) -- ENSEMBLE is the series the published PDF prints.

Checks: male + female = total and the three age groups sum to the total,
within 2 persons (independent rounding), every year.
"""
from __future__ import annotations

import re

import openpyxl
import pandas as pd

# label (start of the "Indicateur" cell) -> (measure, sex, age_group, unit)
_ROWS = {
    "Population de la Côte d'Ivoire": ("count", "total", "Total", "persons"),
    "Population homme": ("count", "male", "Total", "persons"),
    "Population femme": ("count", "female", "Total", "persons"),
    "Population de 0-14 ans": ("count", "total", "0-14", "persons"),
    "Population de 15-59 ans": ("count", "total", "15-59", "persons"),
    "Population de 60 ans et plus": ("count", "total", "60+", "persons"),
    "Taux de croissance démographique": ("growth_rate", "total", "Total", "percent"),
    "Densité de la population": ("density", "total", "Total", "per_km2"),
    "Rapport de masculinité": ("sex_ratio", "total", "Total", "ratio"),
    "Indice synthétique de fécondité (ISF)": ("fertility_rate", "total", "Total", "births_per_woman"),
    "Espérance de vie à la naissance (Total)": ("life_expectancy", "total", "Total", "years"),
    "Espérance de vie à la naissance (hommes)": ("life_expectancy", "male", "Total", "years"),
    "Espérance de vie à la naissance (femmes)": ("life_expectancy", "female", "Total", "years"),
}
_CENSUS_2021 = {"total": 29389150.0, "male": 15344990.0, "female": 14044160.0}


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def _decimals(fmt: str) -> int | None:
    first = fmt.split(";")[0]
    if "0" not in first:
        return None                      # General: shown as stored
    m = re.search(r"0\.(0+)", first)
    return len(m.group(1)) if m else 0


def parse(local_path: str, extras: list[str] | None = None) -> pd.DataFrame:
    wb = openpyxl.load_workbook(local_path, data_only=True)
    if "ENSEMBLE" not in wb.sheetnames:
        raise ValueError(f"ANStat CIV: no ENSEMBLE sheet in {wb.sheetnames}")
    ws = wb["ENSEMBLE"]
    rows = list(ws.iter_rows())
    head = next(r for r in rows if _norm(r[1].value) == "Domaine")
    years = {c.column: str(c.value) for c in head if isinstance(c.value, int)}
    if sorted(years.values()) != [str(y) for y in range(2015, 2026)]:
        raise ValueError(f"ANStat CIV: years {sorted(years.values())}")
    start = next(i for i, r in enumerate(rows) if _norm(r[1].value).startswith("DÉMOGRAPHIE"))
    if "(p)" not in _norm(rows[start][1].value):
        raise ValueError("ANStat CIV: the DÉMOGRAPHIE block is no longer flagged (p)")

    found: dict[str, dict[str, float]] = {}
    for r in rows[start:]:
        if r[1].value and r is not rows[start]:          # next Domaine block
            break
        label = _norm(r[2].value)
        key = next((k for k in _ROWS if _norm(k) == label), None)
        if key is None:
            continue
        vals = {}
        for c in r:
            if c.column in years and c.value is not None:
                if not isinstance(c.value, (int, float)):
                    raise ValueError(f"ANStat CIV {key} {years[c.column]}: {c.value!r}")
                d = _decimals(c.number_format)
                vals[years[c.column]] = float(round(c.value, d)) if d is not None else float(c.value)
        found[key] = vals
    missing = set(_ROWS) - set(found)
    if missing:
        raise ValueError(f"ANStat CIV: rows not found {sorted(missing)}")

    pop = found["Population de la Côte d'Ivoire"]
    for y in pop:
        m, f = found["Population homme"][y], found["Population femme"][y]
        if abs(m + f - pop[y]) > 2:
            raise ValueError(f"ANStat CIV {y}: M + F {m + f} != {pop[y]}")
        a = sum(found[k][y] for k in ("Population de 0-14 ans", "Population de 15-59 ans",
                                      "Population de 60 ans et plus"))
        if abs(a - pop[y]) > 2:
            raise ValueError(f"ANStat CIV {y}: age groups {a} != {pop[y]}")
    for sex, k in (("total", "Population de la Côte d'Ivoire"), ("male", "Population homme"),
                   ("female", "Population femme")):
        if found[k]["2021"] != _CENSUS_2021[sex]:
            raise ValueError(f"ANStat CIV 2021 {sex}: {found[k]['2021']} != census {_CENSUS_2021[sex]}")

    out = []
    for key, vals in found.items():
        measure, sex, age, unit = _ROWS[key]
        for y, v in vals.items():
            if y == "2021":
                continue
            out.append({"series_type": "projection", "sex": sex, "age_group": age,
                        "geography": "Total country", "period": y,
                        "frequency": "annual", "measure": measure, "value": v,
                        "unit": unit, "series_code": "CI_ANSTAT_CHIFFRES_CLES_2025"})
    return pd.DataFrame(out)
