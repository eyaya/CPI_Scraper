"""Parser for the INSTAT Mali quarterly GDP (PIB) note PDF (Tier 3, French).

Clean francophone note. Four level tables give the quarterly GDP breakdown; each
row is a branch (or expenditure component) and the columns are the quarters
('Tn_YYYY'). The tables are pure levels (growth is in separate 'Taux de
croissance' tables, which we skip):
  Tableau 1  Valeurs ajoutées du PIB en volume chaîné      -> production constant
  Tableau 3  Valeurs ajoutées du PIB à prix courant        -> production current
  Tableau 4  composantes du PIB … optique de dépenses en volume chaîné -> exp constant
  Tableau 6  composantes du PIB … optique de dépenses aux prix courants -> exp current

Per row we read the label and the French numbers in order ('2 494,4' / '2494,4');
the first len(periods) values are the quarterly levels. Values are billion XOF
(FCFA). Everything as published; nothing derived.
"""
from __future__ import annotations
import re
import pdfplumber
import pandas as pd

_PERIOD_RE = re.compile(r"T([1-4])[_ ]?(20\d\d)")

_OUT_COLS = ["approach", "category", "category_group", "series_code", "geography",
             "period", "frequency", "price_basis", "seasonal_adjustment",
             "measure", "value", "unit", "base_period"]


def _spec(caption: str):
    c = caption.lower()
    if "taux de croissance" in c:
        return None
    approach = "expenditure" if "composante" in c else "production"
    if "prix courant" in c:
        basis = "current"
    elif "volume" in c or "chaîn" in c or "chain" in c:
        basis = "constant"
    else:
        return None
    return approach, basis


def _caption(page) -> str:
    for line in (page.extract_text() or "").split("\n"):
        if re.search(r"tableau\s*\d", line, re.I):
            return line
    return ""


def _frval(cell):
    s = str(cell or "").strip().replace(" ", "").replace(" ", "").replace(",", ".")
    return float(s) if re.fullmatch(r"-?\d+(\.\d+)?", s) else None


def _periods(tb):
    for row in tb[:4]:
        pers, seen = [], set()
        for c in row:
            m = _PERIOD_RE.search(re.sub(r"\s", "", str(c or "")))
            if m:
                p = f"{m.group(2)}-Q{m.group(1)}"
                if p not in seen:
                    seen.add(p); pers.append(p)
        if len(pers) >= 4:
            return pers
    return []


def _parse_table(tb, approach, basis):
    periods = _periods(tb)
    if not periods:
        return []
    n = len(periods)
    base = "chained volume" if basis == "constant" else ""
    rows = []
    for row in tb:
        cells = [("" if c is None else str(c).strip()) for c in row]
        label = next((re.sub(r"\s+", " ", c).strip() for c in cells
                      if c and re.search(r"[A-Za-zÀ-ÿ]{3,}", c) and _frval(c) is None), "")
        if not label or len(label) < 3 or _PERIOD_RE.search(label) \
                or any(k in label.lower() for k in ("libellé", "branche", "tableau")):
            continue
        nums = [v for v in (_frval(c) for c in cells) if v is not None]
        if len(nums) < n:
            continue
        low = label.lower()
        ap = "aggregate" if low == "pib" or "produit intérieur brut" in low else approach
        for period, v in zip(periods, nums[0:n]):
            rows.append({
                "approach": ap, "category": label, "category_group": "",
                "series_code": "", "geography": "National", "period": period,
                "frequency": "quarterly", "price_basis": basis,
                "seasonal_adjustment": "nsa", "measure": "level",
                "value": v, "unit": "XOF billion",
                "base_period": base if basis == "constant" else "",
            })
    return rows


# --- the back-series workbook (tab-pib*_pub.xlsx) ------------------------------
# INSTAT posts, beside each note, a workbook with the whole quarterly series from
# T1_2004: production by branch and expenditure by component, in chained volume
# (reference 2015), at current prices and at previous-year prices, each with its
# y/y growth sheet, plus implicit deflators (2015 = 100) and their y/y change.
# sheet -> (approach, price_basis, measure, unit, base_period, what)
# SKIPPED: the two 'N/N-1 Variation' sheets. Their columns run one quarter
# behind their own level sheet (the value headed T2 2015, 8.09, is the y/y
# ratio of the previous-year-price levels of T1 2015/T1 2014), so which
# quarter a figure belongs to cannot be read off the header.
_WB_SHEETS = {
    "PIB_Trim_CHainé_Millards_Fcfa": ("production", "constant", "level", "XOF billion", "chained volume (reference 2015)", ""),
    "Taux croissance_PIB_Trim_CHainé": ("production", "constant", "growth_yoy", "percent", "chained volume (reference 2015)", ""),
    "PIB_Trim_CRT_Milliards_FCFA": ("production", "current", "level", "XOF billion", "", ""),
    "PIB_Trim_CRT_TAUX_CROISANCE": ("production", "current", "growth_yoy", "percent", "", ""),
    "TPub_PIB_Trim_N_N-1_Millards": ("production", "constant", "level", "XOF billion", "previous year prices", ""),
    "Deflateur_Valeur": ("production", "not_applicable", "deflator", "index", "2015=100", ""),
    "Deflateur_Variation": ("production", "not_applicable", "growth_yoy", "percent", "", "deflator"),
    "PIB_emploi_trim_Chain_Milliards": ("expenditure", "constant", "level", "XOF billion", "chained volume (reference 2015)", ""),
    "PIB_emploi_trim_Chain_Variation": ("expenditure", "constant", "growth_yoy", "percent", "chained volume (reference 2015)", ""),
    "PIB_emploi_trim_CRT_Milliards": ("expenditure", "current", "level", "XOF billion", "", ""),
    "PIB_emploi_trim_CRT_Variation": ("expenditure", "current", "growth_yoy", "percent", "", ""),
    "PIB_emploi_trim_N_N-1_Milliards": ("expenditure", "constant", "level", "XOF billion", "previous year prices", ""),
    "Deflateur_trimestriel_Milliars": ("expenditure", "not_applicable", "deflator", "index", "2015=100", ""),
    "Deflateur_trimestriel_Variation": ("expenditure", "not_applicable", "growth_yoy", "percent", "", "deflator"),
}
_WB_RE = re.compile(r"tab-pib.*?(\d)t(20\d\d)", re.I)


def _wb_edition(path: str) -> str:
    import os
    m = _WB_RE.search(os.path.basename(path))
    return f"T{m.group(1)} {m.group(2)}" if m else os.path.basename(path)


def _parse_workbook(path: str) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    edition = _wb_edition(path)
    rows = []
    for sheet, (approach, basis, measure, unit, base, what) in _WB_SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        grid = list(wb[sheet].iter_rows(values_only=True))
        hi = next((i for i, r in enumerate(grid)
                   if any(isinstance(v, str) and _PERIOD_RE.fullmatch(v.strip()) for v in r)), None)
        if hi is None:
            continue
        cols = [(j, _PERIOD_RE.fullmatch(v.strip())) for j, v in enumerate(grid[hi])
                if isinstance(v, str) and _PERIOD_RE.fullmatch(v.strip())]
        periods = [(j, f"{m.group(2)}-Q{m.group(1)}") for j, m in cols]
        # DEFECT in the T2 2026 workbook: Deflateur_Valeur's last column is
        # headed 'T1_2027' (it holds T2 2026: PIB 147.63 = current 5782.0 /
        # chained 3916.5). Header out of sequence -> the column is skipped,
        # not relabelled.
        ok, prev = [], None
        for j, p in periods:
            if prev is not None and p <= prev:
                break
            if prev is not None and (int(p[:4]) * 4 + int(p[-1])) - (int(prev[:4]) * 4 + int(prev[-1])) != 1:
                break
            ok.append((j, p)); prev = p
        first = ok[0][0]
        group = f"INSTAT back-series workbook ({edition})"
        if what:
            group += ", deflator y/y % change"
        for r in grid[hi + 1:]:
            label = next((re.sub(r"\s+", " ", v).strip() for v in r[:first]
                          if isinstance(v, str) and v.strip()), "")
            if not label or label.lower().startswith("source"):
                continue
            ap = "aggregate" if label == "PIB" and approach == "production" else approach
            for j, period in ok:
                v = r[j] if j < len(r) else None
                if not isinstance(v, (int, float)):
                    continue
                rows.append({
                    "approach": ap, "category": label, "category_group": group,
                    "series_code": "", "geography": "National", "period": period,
                    "frequency": "quarterly", "price_basis": basis,
                    "seasonal_adjustment": "nsa", "measure": measure,
                    "value": float(v), "unit": unit, "base_period": base,
                })
    return rows


def parse(pdf_path: str, extras: list[str] | None = None) -> pd.DataFrame:
    rows = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            spec = _spec(_caption(page))
            if not spec:
                continue
            tbls = [tb for tb in page.extract_tables() if len(tb) >= 8]
            for tb in tbls:
                rows.extend(_parse_table(tb, *spec))
    if not rows:
        raise ValueError("no GDP rows parsed from INSTAT Mali note")
    df = pd.DataFrame.from_records(rows)[_OUT_COLS].drop_duplicates(
        ["approach", "category", "period", "price_basis"])
    # The workbook supplies only what the note does not print: its LEVELS for
    # quarters outside the note's window (the note's own levels win there), and
    # every growth / deflator row (the note's parse skips those tables).
    note_periods = set(df["period"])
    wb_rows = []
    for x in extras or []:
        if _WB_RE.search(x) and x.lower().endswith(".xlsx"):
            wb_rows.extend(_parse_workbook(x))
    if wb_rows:
        wb = pd.DataFrame.from_records(wb_rows)[_OUT_COLS]
        drop = ((wb["measure"] == "level") & (wb["base_period"] != "previous year prices")
                & wb["period"].isin(note_periods))
        df = pd.concat([df, wb[~drop]], ignore_index=True)
    return df
