"""Back series for ANSD Senegal GDP (helpers for senegal_ansd_gdp, read as extras).

Two ANSD workbooks, each a different vintage or base, each tagged in
category_group and never chained to another:

* 'Tableaux-publication_CN2025_provisoire_...xlsx' -- the renovated accounts
  (BASE 2021) retropolated to 2014 and extended to 2025 (provisional):
    'PIB_optique_production'       value added / GDP by ISIC activity at current
                                   prices, in chained volume (reference 2021)
                                   and the annual volume change (%);
    'Ressources et emplois de B&S' the goods-and-services balance (GDP, imports,
                                   consumption, GFCF, stocks, exports) in the
                                   same three forms.
  Million FCFA. Rows for the years the primary workbook prints are dropped by
  the caller (that edition's own figures stand); this supplies 2014-2020 and
  2025.

* 'CN_Edition_1999_2023_vf.xlsx' -- the PREVIOUS base (BASE 2014), 1999-2023
  (2023 semi-final, starred): the goods-and-services account and GDP by branch,
  at current prices and in volume (reference 2014). Billion FCFA. A separate
  base, emitted whole.

Everything as published; nothing derived.
"""
from __future__ import annotations
import os
import re
import openpyxl

_YEAR = re.compile(r"^((?:19|20)\d\d)\*?$")


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\xa0", "").replace(" ", "")
    return float(s) if re.fullmatch(r"-?\d+(\.\d+)?([eE]-?\d+)?", s) else None


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def _year_cols(ws, r):
    out = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        m = _YEAR.match(str(int(v)) if isinstance(v, (int, float)) and not isinstance(v, bool)
                        else _norm(v))
        if m:
            out[c] = m.group(1)
    return out if len(out) >= 3 else {}


def _row(approach, label, code, group, year, basis, measure, value, unit, base):
    return {"approach": approach, "category": label, "category_group": group,
            "series_code": code, "geography": "National", "period": year,
            "frequency": "annual", "price_basis": basis,
            "seasonal_adjustment": "nsa", "measure": measure, "value": value,
            "unit": unit if measure == "level" else "percent",
            "base_period": base if basis == "constant" else ""}


def _block_spec(title: str):
    t = title.lower()
    if "volume" in t and ("%" in t or "evolution" in t or "évolution" in t
                          or "variation" in t or "croissance" in t):
        return "constant", "growth_yoy"
    if "volume" in t or "constant" in t or "chaîn" in t or "chain" in t:
        return "constant", "level"
    if "courant" in t:
        return "current", "level"
    return None


def provisional_2025(path: str) -> list[dict]:
    """The base-2021 workbook 2014-2025 (see module docstring)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    tag = "ANSD comptes nationaux base 2021, retropolated 2014-2020 + 2025 provisional (MAJ 1, 22-09-2026)"
    base = "chained volume, reference 2021"
    rows = []
    for name in wb.sheetnames:
        low = name.lower()
        if "production" in low:
            sheet_ap = "production"
        elif "ressources" in low:
            sheet_ap = "expenditure"
        else:
            continue
        ws = wb[name]
        title, spec, ycols, section = "", None, {}, ""
        for r in range(1, ws.max_row + 1):
            yc = _year_cols(ws, r)
            if yc:
                spec, ycols, section = _block_spec(title), yc, ""
                continue
            texts = [_norm(ws.cell(r, c).value) for c in range(1, 4)]
            nums = {y: _num(ws.cell(r, c).value) for c, y in ycols.items()}
            has = any(v is not None for v in nums.values())
            if not has:
                t = next((x for x in texts if x), "")
                if texts[2] in ("Ressources", "Emplois"):
                    section = texts[2]
                elif len(t) > 30 and not t.lower().startswith(("source", "(*")):
                    title, spec, ycols = t, None, {}     # a new block's title
                continue
            if spec is None:
                continue
            code, label = texts[1], texts[2]
            if not label:
                continue
            basis, measure = spec
            ap = sheet_ap
            if sheet_ap == "production" and label == "PIB":
                ap = "aggregate"
            grp = f"{tag}, {section}" if section else tag
            for c, y in ycols.items():
                v = nums[y]
                if v is None:
                    continue
                if measure == "growth_yoy" and "%" in (ws.cell(r, c).number_format or ""):
                    v *= 100.0
                rows.append(_row(ap, label, code, grp, y, basis, measure, v,
                                 "million FCFA", base))
    wb.close()
    return rows


def base2014_edition(path: str) -> list[dict]:
    """The previous base (2014), 1999-2023 (see module docstring)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    tag = "ANSD comptes nationaux BASE 2014 (Edition 1999-2023; 2023 semi-final)"
    rows = []
    ycols, spec, ap = {}, None, "expenditure"
    for r in range(1, ws.max_row + 1):
        head = _norm(ws.cell(r, 2).value)
        if "COMPTE DE BIENS" in head.upper():
            ap = "goods_services_account"
        elif "PAR BRANCHE" in head.upper():
            ap = "production"
        yc = _year_cols(ws, r)
        if yc:
            ycols = yc
            t = head.lower()
            if "volume" in t:
                spec = ("constant", "level", "volume, reference prices of 2014")
            elif "courant" in t:
                spec = ("current", "level", "")
            else:
                spec = None
            continue
        if not ycols or spec is None or not head:
            continue
        basis, measure, base = spec
        label = head
        m = re.match(r"^([A-Z]{1,2}\d{2}|Z[AB]0)\s+(.*)$", label)
        code = m.group(1) if m else ""
        upper = label.upper()
        if ap == "goods_services_account":
            approach = "expenditure"
        else:
            approach = "aggregate" if "PRODUIT INTERIEUR BRUT" in upper else "production"
        for c, y in ycols.items():
            v = _num(ws.cell(r, c).value)
            if v is None:
                continue
            rows.append(_row(approach, label, code, tag, y, basis, measure, v,
                             "billion FCFA", base))
    wb.close()
    return rows


def read(extras: list[str] | None) -> tuple[list[dict], list[dict]]:
    """-> (base-2021 back-series rows, base-2014 rows) from whichever extras exist."""
    b21, b14 = [], []
    for x in extras or []:
        n = os.path.basename(x).lower()
        if "cn2025_provisoire" in n:
            b21 += provisional_2025(x)
        elif n.startswith("cn_edition_1999"):
            b14 += base2014_edition(x)
    return b21, b14
