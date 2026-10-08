"""Lesotho — BOS Quarterly GDP publication tables (Tier 2 workbook).

FOUND BEHIND A JAVASCRIPT CATALOGUE. bos.gov.ls builds its publication list in
the browser from a `catalogueData` array inside publications.htm, so crawlers
saw only policy PDFs and Lesotho was recorded as having no GDP file. The array
lists "QGDP Publication Tables <n> Quarter <year>" under
/Bos_Reports/Copy%20of%20Economics/, each a ZIP holding one workbook.

THE WORKBOOK -- production approach, 2007 onward, four sheets:

* GDP_CP / GDP_KP -- levels at CURRENT and CONSTANT 2012 prices, million
  maloti (the report's Tables 2-3: "GDP at Constant Prices, Million Maloti";
  "QGDP at current and constant 2012 prices"). Annual rows (Year, no quarter)
  then quarterly rows (Year, Qn). Columns: fifteen industry groups under
  ISIC letters (A, B, C, DE, F ... RST), FISIM, all industries at basic
  prices, taxes on products, GDP, and seasonally adjusted GDP.
* GDP_%Changes -- growth. VERIFIED against the level sheets, not assumed from
  the sheet name: every column is REAL YEAR-ON-YEAR growth (2026 Q1 GDP 2.822 =
  KP 2026Q1 / KP 2025Q1; annual 2025 3.874 = KP 2025 / KP 2024) EXCEPT the
  seasonally adjusted GDP column, which is QUARTER-ON-QUARTER growth of
  adjusted real GDP (-0.766 = KP_sa 2026Q1 / KP_sa 2025Q4).
* Contributions -- despite the name, SHARES of GDP at CURRENT prices: each
  quarter's column sums to 100 and equals CP industry / CP GDP exactly (2019 Q1
  agriculture 4.5235). Filed as `share`, not `contribution` (which in this
  schema means percentage points of growth).

CHECKS RUN ON EVERY PARSE: annual levels equal the sum of their four quarters
where all four are published (2025's stale annual components are dropped --
see _STALE_ANNUAL); each quarter's shares sum to 100; GDP = all
industries at basic prices + taxes on products.

The industry labels are printed hyphenated for column width ("Agricul-ture,
forestry, fishing"); the hyphenation is closed up, the words are not changed.

CROSS-CHECK: real GDP 2025 = LSL 22,632.8 million (2012 prices), growth 3.87%;
2026 Q1 real GDP growth 2.82% y/y, seasonally adjusted -0.77% q/q.
"""
from __future__ import annotations

import os
import re
import zipfile

import openpyxl
import pandas as pd
import pdfplumber

UNIT = "LSL million"
BASE = "Constant 2012 prices"
_AGG = {"All industries at basic prices", "Taxes on products", "GDP"}

# A STALE ANNUAL ROW. In the Q1-2026 edition every year 2007-2024 is internally
# consistent, but the 2025 ANNUAL row mixes two vintages: its GDP column is
# the sum of the four 2025 quarters (current vintage, growth 3.874%), while its
# industry, GVA and tax columns are an older, lower estimate that no longer
# adds up to that GDP (current prices 1,080 short; constant 693 short), and
# the annual growth rates beside them use the old vintage too --
# manufacturing 2.2% on the annual row against 9.1% from the same file's four
# quarters. Publishing both would put two contradictory 2025s in one series.
# So for 2025 only the annual GDP is kept; the component rows are dropped, and
# the four 2025 quarters carry the year. Checked every run: if BOS brings the
# annual row back into line (GVA = sum of quarters), the parser raises so the
# exclusion can be lifted.
_STALE_ANNUAL = {"2025"}


def _clean(text) -> str:
    s = re.sub(r"\s+", " ", str(text or "")).strip()
    # Column-width hyphenation: "Agricul-ture", "Construc-tion", "indu-stries"
    return re.sub(r"(?<=[a-z])-(?=[a-z])", "", s)


def _grid(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    codes, heads = rows[0], rows[1]
    cols = []
    for j in range(2, len(heads)):
        if heads[j] is None:
            continue
        name = _clean(heads[j])
        if name == "GDP Seasonally adjusted":
            cols.append((j, "GDP", "GDP", "saa"))
        else:
            code = "" if codes[j] is None else str(codes[j]).strip()
            cols.append((j, name, code, "nsa"))
    data = []
    for r in rows[2:]:
        if r[0] is None:
            continue
        year = str(r[0]).strip()
        if not re.fullmatch(r"\d{4}", year):
            continue
        q = str(r[1]).strip() if r[1] is not None else ""
        if q and not re.fullmatch(r"Q[1-4]", q):
            raise ValueError(f"{ws.title}: unreadable quarter {r[1]!r}")
        data.append((f"{year}-{q}" if q else year, "quarterly" if q else "annual", r))
    return cols, data


def _stale(period: str, cat: str) -> bool:
    return period in _STALE_ANNUAL and cat != "GDP"


def _check_still_stale(ws) -> None:
    cols, data = _grid(ws)
    gva = {p: r[j] for p, f, r in data for j, c, code, sa in cols
           if c == "All industries at basic prices" and sa == "nsa"}
    for year in _STALE_ANNUAL:
        qs = [gva.get(f"{year}-Q{i}") for i in range(1, 5)]
        if year in gva and None not in qs and abs(sum(qs) - gva[year]) <= 0.5:
            raise ValueError(f"{ws.title}: the {year} annual row now matches its "
                             f"quarters -- remove {year} from _STALE_ANNUAL")


def _row(cat, code, period, freq, basis, sa, measure, value, unit, base):
    return {"approach": "aggregate" if cat in _AGG else "production",
            "category": cat, "category_group": "", "series_code": code,
            "geography": "National", "period": period, "frequency": freq,
            "price_basis": basis, "seasonal_adjustment": sa,
            "measure": measure, "value": float(value), "unit": unit,
            "base_period": base}


def _levels(ws, basis: str) -> list[dict]:
    cols, data = _grid(ws)
    out = []
    for period, freq, r in data:
        for j, cat, code, sa in cols:
            if isinstance(r[j], (int, float)) and not _stale(period, cat):
                out.append(_row(cat, code, period, freq, basis, sa, "level",
                                r[j], UNIT, BASE if basis == "constant" else ""))
    # annual = sum of its four quarters, and GDP = GVA + taxes
    idx = {(o["period"], o["category"], o["seasonal_adjustment"]): o["value"]
           for o in out}
    for (period, cat, sa), v in list(idx.items()):
        if re.fullmatch(r"\d{4}", period) and sa == "nsa":
            qs = [idx.get((f"{period}-Q{i}", cat, sa)) for i in range(1, 5)]
            if None not in qs and abs(sum(qs) - v) > 0.5:
                raise ValueError(f"{ws.title} {cat} {period}: annual {v} != "
                                 f"sum of quarters {sum(qs)}")
        if cat == "GDP" and sa == "nsa":
            gva = idx.get((period, "All industries at basic prices", sa))
            tax = idx.get((period, "Taxes on products", sa))
            if gva is not None and tax is not None and abs(gva + tax - v) > 0.5:
                raise ValueError(f"{ws.title} {period}: GDP != GVA + taxes")
    return out


def _growth(ws) -> list[dict]:
    cols, data = _grid(ws)
    out = []
    for period, freq, r in data:
        for j, cat, code, sa in cols:
            if isinstance(r[j], (int, float)) and not _stale(period, cat):
                measure = "growth_qoq" if sa == "saa" else "growth_yoy"
                out.append(_row(cat, code, period, freq, "constant", sa, measure,
                                r[j], "percent", ""))
    return out


def _shares(ws) -> list[dict]:
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    years, quarters = rows[0], rows[1]
    periods, year = [], None
    for j in range(1, len(quarters)):
        if years[j] is not None:
            year = str(int(years[j]))
        if quarters[j] and re.fullmatch(r"Q[1-4]", str(quarters[j])):
            periods.append((j, f"{year}-{quarters[j]}"))
    out, sums = [], {}
    for r in rows[2:]:
        if r[0] is None:
            continue
        cat = _clean(r[0])
        for j, period in periods:
            if isinstance(r[j], (int, float)):
                if cat == "Total":
                    sums.setdefault(period, [0.0, None])[1] = r[j]
                    continue
                sums.setdefault(period, [0.0, None])[0] += r[j]
                out.append(_row(cat, "", period, "quarterly", "current", "nsa",
                                "share", r[j], "percent", ""))
    for period, (s, total) in sums.items():
        if abs(s - 100) > 0.2:
            raise ValueError(f"Contributions {period}: shares sum to {s:.2f}")
    return out


# ------------------------------------------------ annual expenditure tables
# EXTRA: the "Annual National Accounts" release (PDF inside a ZIP). Only its
# EXPENDITURE tables are read -- the approach the quarterly workbook lacks.
# Its production tables are a different vintage from the quarterly workbook
# and are not mixed in. "GDP at market prices" repeats "Expenditure on GDP"
# and is skipped, so nothing here collides with the workbook's GDP.
_EXP_TABLES = [  # caption, measure, price basis, unit, base
    (r"^Table 3a: Expenditure on GDP, current prices", "level", "current", UNIT, ""),
    (r"^Table 3b: Expenditure on GDP, per cent shares", "share", "current", "percent", ""),
    (r"^Table 4a: Expenditure on GDP, constant prices", "level", "constant", UNIT, BASE),
    (r"^Table 4b: Expenditure on GDP, per cent annual changes", "growth_yoy", "constant", "percent", ""),
    (r"^Table 5: Deflators, 2012 = 100", "deflator", "not_applicable", "index", "2012 = 100"),
]
_EXP_ROWS = [
    "Final consumption expenditure, government",
    "Final consumption expenditure, households",
    "Final consumption expenditure, NPISH",
    "Gross fixed capital formation", "Changes in inventories",
    "Gross domestic expenditure", "Exports of goods and services",
    "Less: Imports of goods and services", "Expenditure on GDP",
]
_NUM = r"-?\s?\d{1,3}(?:,\d{3})*(?:\.\d+)?"


def _annual_pdf(path: str) -> str:
    if path.lower().endswith(".zip"):
        with zipfile.ZipFile(path) as z:
            pdfs = [m for m in z.namelist() if m.lower().endswith(".pdf")]
            if len(pdfs) != 1:
                raise ValueError(f"{path}: expected one PDF, found {pdfs}")
            z.extract(pdfs[0], os.path.dirname(path))
            return os.path.join(os.path.dirname(path), pdfs[0])
    return path


def _expenditure(path: str) -> list[dict]:
    with pdfplumber.open(_annual_pdf(path)) as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
    out = []
    for cap, measure, basis, unit, base in _EXP_TABLES:
        m = re.search(cap + r"(?![^\n]*\.{4})", text, re.M)
        if not m:
            raise ValueError(f"annual accounts: {cap!r} not found")
        lines = text[m.end():].splitlines()
        years = re.findall(r"\b(20\d\d)\b", lines[1])
        if len(years) < 5:
            raise ValueError(f"{cap}: no year header")
        got = {}
        for ln in lines[2:2 + len(_EXP_ROWS) + 2]:
            lab = next((r for r in _EXP_ROWS if ln.startswith(r)), None)
            if lab is None:
                continue
            nums = re.findall(_NUM, ln[len(lab):])
            if not nums:
                continue          # Table 4b prints no growth for inventories
            if len(nums) != len(years):
                raise ValueError(f"{cap} {lab}: {len(nums)} values for {len(years)} years")
            got[lab] = [float(n.replace(",", "").replace(" ", "")) for n in nums]
        if measure == "level":
            for k, y in enumerate(years):
                g = got
                gde = sum(g[r][k] for r in _EXP_ROWS[:5])
                if abs(gde - g["Gross domestic expenditure"][k]) > 3:
                    raise ValueError(f"{cap} {y}: components != domestic expenditure")
                gdp = (g["Gross domestic expenditure"][k] + g["Exports of goods and services"][k]
                       - g["Less: Imports of goods and services"][k])
                if abs(gdp - g["Expenditure on GDP"][k]) > 3:
                    raise ValueError(f"{cap} {y}: GDE + X - M != GDP")
        for lab, vals in got.items():
            for y, v in zip(years, vals):
                out.append({"approach": "expenditure", "category": lab,
                            "category_group": "", "series_code": "",
                            "geography": "National", "period": y,
                            "frequency": "annual", "price_basis": basis,
                            "seasonal_adjustment": "not_applicable" if measure == "deflator" else "nsa",
                            "measure": measure, "value": v, "unit": unit,
                            "base_period": base})
    return out


def parse(local_path: str, extras: list[str] | None = None) -> pd.DataFrame:
    wb = openpyxl.load_workbook(local_path, data_only=True)
    need = {"GDP_CP", "GDP_KP", "GDP_%Changes", "Contributions"}
    if not need <= set(wb.sheetnames):
        raise ValueError(f"{os.path.basename(local_path)}: sheets {wb.sheetnames} "
                         f"lack {sorted(need - set(wb.sheetnames))}")
    _check_still_stale(wb["GDP_CP"])
    rows = (_levels(wb["GDP_CP"], "current") + _levels(wb["GDP_KP"], "constant")
            + _growth(wb["GDP_%Changes"]) + _shares(wb["Contributions"]))
    for ex in extras or []:
        rows += _expenditure(ex)
    return pd.DataFrame.from_records(rows)
