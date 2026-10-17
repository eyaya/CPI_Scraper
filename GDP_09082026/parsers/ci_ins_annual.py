"""Annual back series for Côte d'Ivoire: the former INS's own national-accounts
workbooks (ins.ci, now offline; fetched as Wayback `id_` copies of the INS
files). Not a registry parser: `ci_anstat_gdp.parse` hands it the `.xlsx`
extras.

Three workbooks, three vintages -- none chained, each kept under its own
`category_group` and base:

* `SYNTHESE DES COMPTES NATIONAUX 1996_2015 RETROPOLES.xlsx` -- 1996-2015, two
  series side by side on each sheet: the retropolated accounts (SNA 2008, base
  2015) and the earlier definitive accounts (SNA 1993, base 1996). Sheets
  Courant (current, XOF billion), Volume (XOF billion) and
  Taux_de_croissance_reelle. The 'Resume' sheet repeats the growth rates and
  the comparison block on 'Courant' is a ratio of the two series; neither is
  read.
* `compte_nationaux_2021CD_2022CP.xlsx` -- 2015-2022 (CD = comptes définitifs,
  CP = provisoires), SNA 2008: current prices, prices of the previous year,
  real growth, share of GDP, price change, contribution to growth.
* `ComptesNationauxAnnuel_prix2009.xlsx` -- 1996-2016 value added by branch and
  GDP by expenditure at constant 2009 prices. The workbook does not print its
  scale; the figures are in XOF million (PIB 1996 = 9,493,239 against the INS
  synthesis' current-price 6,209.8 billion under the 1993 base), and the unit
  says so.

Percent cells are stored as fractions with a percent number format; they are
read as the percent the workbook displays (the burkina/gambia/mauritania
convention). Everything as published; nothing derived.
"""
from __future__ import annotations
import os
import re
import openpyxl

_YEAR = re.compile(r"^\s*((?:19|20)\d\d)\s*(?:CD|CP)?\s*$", re.I)
_AGG = {"PIB", "PIB AU COÛT DES FACTEURS", "VALEUR AJOUTEE"}
_EXP_START = ("IMPORTATIONS", "CONSOMMATION FINALE")
_HEADS = ("CONSOMMATION FINALE", "FBCF", "FORMATION BRUTE DE CAPITAL FIXE",
          "EXPORTATIONS", "IMPORTATIONS")


def _clean(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip()


def _rec(approach, cat, group, year, basis, measure, value, unit, base):
    return {"approach": approach, "category": cat, "category_group": group,
            "series_code": "", "geography": "National", "period": year,
            "frequency": "annual", "price_basis": basis,
            "seasonal_adjustment": "not_applicable", "measure": measure,
            "value": float(value), "unit": unit, "base_period": base}


def _read_grid(ws, label_col, group_of, spec_of):
    """Generic reader: a year header row, then labelled rows. `group_of(row)`
    returns the block's (group, base) when a row opens a new block (or None);
    `spec_of` gives (measure, basis, unit)."""
    rows, years, group, base = [], {}, None, ""
    approach, head = "production", ""
    measure, basis, unit = spec_of
    for r in ws.iter_rows():
        cells = {c.column - 1: c for c in r if c.value is not None}
        yr = {j: _YEAR.match(str(c.value)).group(1) for j, c in cells.items()
              if _YEAR.match(str(c.value))}
        if len(yr) >= 4:
            years, approach, head = yr, "production", ""
            continue
        g = group_of(cells)
        if g:
            group, base = g
        lab = _clean(cells[label_col].value) if label_col in cells else ""
        if not lab or not years or not group:
            continue
        if lab.upper().startswith(("EN %", "CD :", "CP :")):
            continue
        U = lab.upper()
        if U.startswith(_EXP_START):
            approach = "expenditure"
        if U in _HEADS or U.startswith("FORMATION BRUTE"):
            head = lab
            cat = lab
        elif approach == "expenditure" and head and U not in (
                "TOTAL EMPLOIS", "TOTAL EMPLOIS FINAL", "TOTAL RESSOURCES",
                "SOLDE EXTERIEUR", "VARIATION DE STOCK", "VARIATION DE STOCKS"):
            cat = f"{head}: {lab}"
        else:
            cat = lab
        appr = "aggregate" if U in _AGG else approach
        for j, y in years.items():
            c = cells.get(j)
            if c is None or not isinstance(c.value, (int, float)):
                continue
            v = c.value * 100 if "%" in (c.number_format or "") else c.value
            rows.append(_rec(appr, cat, group, y, basis, measure, v, unit, base))
    return rows


def _synthese(wb, path):
    out = []
    specs = {"Courant": ("level", "current", "XOF billion"),
             "Volume": ("level", "constant", "XOF billion"),
             "Taux_de_croissance_reelle": ("growth_yoy", "constant", "percent")}
    for name, spec in specs.items():
        if name not in wb.sheetnames:
            continue

        def group_of(cells, name=name, spec=spec):
            a = _clean(cells[0].value) if 0 in cells else ""
            if a.startswith("COMPTES RETROPOLES"):
                base = "volume, SNA 2008 base 2015" if spec[0] == "level" \
                    and spec[1] == "constant" else ""
                return (f"INS synthèse 1996-2015, {name}: comptes rétropolés "
                        f"SCN 2008 base 2015", base)
            if a.startswith("COMPTES DEFINITIFS"):
                base = "volume, SNA 1993 base 1996" if spec[0] == "level" \
                    and spec[1] == "constant" else ""
                return (f"INS synthèse 1996-2015, {name}: comptes définitifs "
                        f"SCN 1993 base 1996", base)
            if a.startswith("VARIATION DES COMPTES"):
                return ("__skip__", "")
            return None
        rows = _read_grid(wb[name], 1, group_of, spec)
        out.extend(r for r in rows if r["category_group"] != "__skip__")
    return out


def _cn_2021_2022(wb, path):
    out = []
    specs = {
        "Tab_Courant": ("level", "current", "XOF billion", ""),
        "Tab_Constant_N-1": ("level", "constant", "XOF billion",
                             "prices of the previous year (N-1), SNA 2008"),
        "Tab_taux de croissance": ("growth_yoy", "constant", "percent", ""),
        "Tab_Poids dans PIB": ("share", "current", "percent", ""),
        "Tab_évolution des prix": ("growth_yoy", "not_applicable", "percent", ""),
        "Tab_Contribution ": ("contribution", "constant", "percentage points", ""),
    }
    for name, (m, b, u, base) in specs.items():
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        title = _clean(ws["A2"].value or ws["A1"].value)
        group = f"INS comptes nationaux 2021CD-2022CP, {title}"
        out.extend(_read_grid(ws, 0, lambda cells, g=(group, base): g, (m, b, u)))
    return out


def _prix2009(wb, path):
    ws = wb.worksheets[0]
    group = ("INS comptes nationaux annuels, " + _clean(ws["A1"].value) +
             " (et PIB par la demande)")
    rows = _read_grid(ws, 0, lambda cells: (group, "constant 2009 prices"),
                      ("level", "constant", "XOF million"))
    return rows


def parse_workbook(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    name = os.path.basename(path).lower()
    if "synthese" in name:
        rows = _synthese(wb, path)
    elif "2021cd" in name:
        rows = _cn_2021_2022(wb, path)
    elif "prix2009" in name:
        rows = _prix2009(wb, path)
    else:
        return []
    if not rows:
        raise ValueError(f"no national-accounts rows read from {path}")
    return rows
