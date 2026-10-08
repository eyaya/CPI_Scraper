"""Gambia Bureau of Statistics (GBoS) CPI — national chained index series.

Parser for the monthly CPI workbook from the GBoS data portal. The
"NChained Link Series to Publish" sheet is the national series: a COICOP column, a
label column, two weight columns, then one column per month (the header carries the
month as a date). Rows are the all-items index and every COICOP division/sub-class.

We emit the published index for each COICOP item and month (measure = index). The
workbook gives index levels only (no inflation rates), so nothing is derived.
"""
from __future__ import annotations
import datetime as dt
import re
import pandas as pd
from openpyxl import load_workbook

_SHEET = "NChained Link Series to Publish"
_CODE_RE = re.compile(r"^\d{1,2}(?:\.\d+)*$")


def _period(v):
    """A month-header cell -> 'YYYY-MM', from a date or a YYYY-MM(-DD) string."""
    if isinstance(v, dt.datetime):
        return f"{v.year:04d}-{v.month:02d}"
    s = str(v).strip()
    m = re.match(r"^(\d{4})[-/](\d{1,2})", s)
    return f"{m.group(1)}-{int(m.group(2)):02d}" if m else None


def _norm_code(c: str) -> str:
    c = c.strip()
    if c in ("0", "00", "0.0"):
        return "00"
    head, _, rest = c.partition(".")
    head = head.zfill(2)
    return head if not rest else f"{head}.{rest}"


def _num(v):
    if v is None or str(v).strip() == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse(xlsx_path: str) -> pd.DataFrame:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[_SHEET]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]

    # header row + COICOP column: the cell (anywhere) equal to "COICOP".
    hi = cc = None
    for i, r in enumerate(rows):
        for c, v in enumerate(r):
            if v is not None and str(v).strip().upper() == "COICOP":
                hi, cc = i, c
                break
        if hi is not None:
            break
    if hi is None:
        raise ValueError("gambia_gbos: header row (COICOP) not found")
    header = rows[hi]
    month_cols = [(c, _period(header[c])) for c in range(len(header))
                  if _period(header[c])]
    if not month_cols:
        raise ValueError("gambia_gbos: no month columns found")

    out = []
    for r in rows[hi + 1:]:
        if not r or cc >= len(r) or r[cc] is None:
            continue
        code = str(r[cc]).strip()
        if not _CODE_RE.match(code):
            continue
        coicop = _norm_code(code)
        label = str(r[cc + 1]).strip() if cc + 1 < len(r) and r[cc + 1] is not None else ""
        for c, period in month_cols:
            val = _num(r[c]) if c < len(r) else None
            if val is None:
                continue
            out.append({
                "coicop_code": coicop, "coicop_label": label,
                "geography": "Total country", "period": period,
                "measure": "index", "value": val, "unit": "Index",
                "base_period": "", "frequency": "monthly",
            })
    df = pd.DataFrame(out)
    if df.empty:
        raise ValueError("gambia_gbos: no rows parsed")
    return df.drop_duplicates(["coicop_code", "period"])
