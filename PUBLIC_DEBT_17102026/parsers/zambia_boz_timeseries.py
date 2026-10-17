"""Zambia -- Bank of Zambia, FORTNIGHTLY TIMESERIES workbook (posted on
boz.zm with the Statistics Fortnightly; e.g.
FORTNIGHTLYTIMESERIESending2October2026_1.xlsx): the MONTHLY blocks of its
"Table 17" (Treasury bills) and "Table 18" (GRZ bonds) sheets, the same tables
the fortnightly PDF prints (parsers/zambia_boz_debt.py), going back further.

Called by zambia_boz_debt.parse for an .xlsx extra; not registered on its own.

WHAT IS READ, outstanding balances only (sales, maturities, settlement values
are flows):
* Table 17, "Monthly Totals" block, January 2009 - May 2017: outstanding 91 /
  182 / 273 / 364 days and total (header letters j, k, l, m, n=j+k+l+m).
  After May 2017 the sheet continues WEEKLY ("Week-ending") -- a weekly stock
  is not a month-end figure and is not read.
* Table 18, monthly block from 2009 to May 2016: outstanding by tenor as the
  sheet heads it ("12 months", "18 months", "24 months", "3 year" ... "15
  year") and TOTAL. Only months whose tenors sum to the printed TOTAL are
  read: in 2009-2011 the sheet prints the tenor columns as 0 beside a total,
  or the total itself as 0, and neither can be told from a real figure -- those
  months are listed and left out. A blank tenor cell (the discontinued 12 and
  18 months) is absent, not 0. The 2004 block above it has a different
  column layout and is not read; the weekly block below is not read.

UNIT: the sheets head both tables "( K' 000 )" (Table 18 also "( K' Million
)"), and the series runs straight through the January 2013 kwacha rebasing
(1 000 old = 1 new) with no 1 000-fold break -- Table 17 total January 2012
6 202 958.84, January 2013 6 832 991.75 -- so BoZ states the pre-2013 months
in rebased-kwacha terms (old K million = new K thousand). Filed as printed,
"ZMW thousand", never converted. Labels and series codes match the PDF
reader's, so the workbook's months join the same series; the tenor "24
months" is kept as printed (the PDF from 2022 prints "2 year").

CHECK: every month's tenors sum to its total (to K 1 thousand).
CROSS-CHECK: Table 17 January 2009 total 3 364 364.086 (91 days 155 751);
Table 18 January 2012 TOTAL 5 215 853.9 (5 year 2 471 214.7).
"""
from __future__ import annotations

import openpyxl
import pandas as pd

_UNIT = "ZMW thousand"
_MONTHS = {m: i for i, m in enumerate(
    ["january", "february", "march", "april", "may", "june", "july",
     "august", "september", "october", "november", "december"], 1)}


def _rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _year(v):
    s = str(v).strip() if v is not None else ""
    return int(s) if s.isdigit() and len(s) == 4 else None


def _num(v):
    return v if isinstance(v, (int, float)) else None


def _block(rows, label_col, cols, comp, code, start_at=0):
    """Read a monthly block: year rows and month-name rows in `label_col`,
    from `start_at` to the first "Week-ending"; `cols` maps label -> column,
    "total" last."""
    out, skipped, year = [], [], None
    for r in rows[start_at:]:
        lab = r[label_col] if len(r) > label_col else None
        s = str(lab).strip() if lab is not None else ""
        if s.startswith("Week"):
            break
        if _year(lab):
            year = _year(lab)
            continue
        if s.lower() not in _MONTHS or year is None:
            continue
        vals = {k: _num(r[c]) if c < len(r) else None for k, c in cols.items()}
        period = f"{year}-{_MONTHS[s.lower()]:02d}"
        # A BLANK tenor cell is a tenor the sheet does not print that month
        # (the discontinued 12/18 months); it is absent, not 0.
        parts = [v for k, v in vals.items() if k != "total" and v is not None]
        tot = vals["total"]
        if not tot or not any(parts) or abs(sum(parts) - tot) > 1.0:
            skipped.append(period)
            continue
        for k, v in vals.items():
            if k != "total" and not v:
                continue           # blank, or a tenor not outstanding that month
            out.append({"component": f"{comp}: {k}", "period": period,
                        "frequency": "monthly", "measure": "stock",
                        "value": float(v), "unit": _UNIT, "series_code": code})
    if skipped:
        print(f"  [zambia] {code} workbook: {len(skipped)} month(s) not checkable "
              f"(tenors 0/blank or total 0), left out: {skipped[0]}..{skipped[-1]}")
    return out


def parse_workbook(path: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Table 17" not in wb.sheetnames or "Table 18" not in wb.sheetnames:
        raise ValueError("BoZ timeseries: Table 17 / Table 18 sheets not found")

    # Table 17: the header-letter row ('j' 'k' 'l' 'm' 'n=j+k+l+m') locates
    # the outstanding columns; 'Monthly Totals' opens the block.
    r17 = _rows(wb["Table 17"])
    hdr = next((r for r in r17 if "n=j+k+l+m" in [str(c).strip() for c in r if c]), None)
    if hdr is None:
        raise ValueError("BoZ timeseries Table 17: header letters not found")
    pos = {str(c).strip(): j for j, c in enumerate(hdr) if c is not None}
    cols17 = {"91 days": pos["j"], "182 days": pos["k"], "273 days": pos["l"],
              "364 days": pos["m"], "total": pos["n=j+k+l+m"]}
    start = next(i for i, r in enumerate(r17)
                 if any(str(c).strip() == "Monthly Totals" for c in r if c))
    lab17 = next(j for j, c in enumerate(r17[start]) if c is not None)
    out = _block(r17, lab17, cols17, "Treasury bills outstanding", "BoZ SF T17", start)

    # Table 18: the stock heads are on the row carrying 'TOTAL '; the monthly
    # block starts at the first year row given as a number (2009) after the
    # 2004 block.
    r18 = _rows(wb["Table 18"])
    hdr = next((r for r in r18 if any(str(c).strip() == "TOTAL" for c in r if c)
                and any(str(c).strip() == "15 year" for c in r if c)), None)
    if hdr is None:
        raise ValueError("BoZ timeseries Table 18: stock heads not found")
    cols18, seen_total = {}, False
    for j, c in enumerate(hdr):
        t = str(c).strip() if c is not None else ""
        if t in ("12 months", "18 months", "24 months", "3 year", "5 year",
                 "7 year", "10 year", "15 year"):
            cols18[t] = j
        elif t == "TOTAL":
            cols18["total"] = j
    if "total" not in cols18 or len(cols18) != 9:
        raise ValueError(f"BoZ timeseries Table 18: heads {sorted(cols18)}")
    start = next(i for i, r in enumerate(r18)
                 if len(r) > 1 and isinstance(r[1], int) and _year(r[1]))
    out += _block(r18, 1, cols18, "GRZ bonds outstanding", "BoZ SF T18", start)

    df = pd.DataFrame(out)
    if df.empty or df.period.str[:4].astype(int).min() > 2009:
        raise ValueError("BoZ timeseries: monthly blocks not read")
    return df
