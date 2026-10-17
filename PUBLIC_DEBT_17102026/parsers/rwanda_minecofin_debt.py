"""Rwanda — Ministry of Finance and Economic Planning (MINECOFIN), Investor
Relations > Debt > "Debt Stock" workbook (Debt_Stock_End_<year>.xlsx), sheet
"Public Debt Website": annual end-of-year series, 2011 to the edition year.

PUBLISHER: MINECOFIN (an accepted official compiler). The National Bank of
Rwanda publishes only PRIVATE-sector external debt, so it is not the source.

COVERAGE AS PUBLISHED: the headline is "Public & Publicly Guaranteed Debt"
(PPG) -- central government external debt + general government domestic debt
+ SOE loans. The workbook also prints private external debt and "PPG +
private"; each row keeps its own label, so the perimeter travels in
`component`.

ROWS READ (explicit list -- the sheet mixes stocks, ratios and flows):

* % of GDP (share_of_gdp): PPG; PPG + private; private; domestic; central
  government external; SOEs; external (public + private).
* PPG debt (FRW billion).
* Central government external debt: foreign currency (USD million), in FRW
  billion; and by creditor (multilaterals, bilaterals, commercial; USD
  million).
* SOE loans (FRW billion).
* General government domestic debt (FRW billion; medium/long and short term)
  and in USD million.

NOT READ: debt service, amortisation and interest (flows); currency and
interest-rate compositions; private external debt breakdowns; GDP.
"Foreign currency (million USD) (Excl. IMF SCF)" is REFUSED: for 2011-2020 it
is LARGER than the series including the IMF facility (2011: 1,544.4 vs 984.7),
which its label makes impossible -- the label or the values are wrong, and the
sheet does not say which. The ratio rows are ratios as published (end-period
exchange rate, per the sheet's heading), never recomputed. "…" = not
available, skipped.

CHECKS: CG external by creditor sums to the CG external total, except 2021
where the creditors exceed it by USD 10.2 m (5,839.35 vs 5,829.17) -- kept as
printed and pinned in `_KNOWN_CREDITOR_GAPS`; domestic
medium/long + short = domestic total (both within 0.5). A listed row that is
missing raises.

CROSS-CHECK (End-2025 edition): PPG 2025 73.63% of GDP / FRW 17,220.36 bn;
CG external 2025 USD 9,504.26 m; domestic 2025 FRW 2,772.87 bn; PPG 2011
20.93%.
"""
from __future__ import annotations

import openpyxl
import pandas as pd

# exact label (stripped) -> (component, measure, unit)
_ROWS = {
    "Public & Publicly Guaranteed Debt (% of GDP)": ("Public & publicly guaranteed debt", "share_of_gdp", "percent"),
    "Public & Publicly Guaranteed Debt +Private debt (% of GDP)": ("Public & publicly guaranteed debt + private debt", "share_of_gdp", "percent"),
    "Private Debt % of GDP": ("Private debt", "share_of_gdp", "percent"),
    "Domestic Debt % of GDP": ("Domestic debt", "share_of_gdp", "percent"),
    "Central Government External Debt (% of GDP)": ("Central government external debt", "share_of_gdp", "percent"),
    "SOE's (Domestic+External) % of GDP": ("SOE debt (domestic + external)", "share_of_gdp", "percent"),
    "External Debt (Public+Private) % of GDP": ("External debt (public + private)", "share_of_gdp", "percent"),
    "Public & Publicly Guaranteed Debt (Billion FRW)": ("Public & publicly guaranteed debt", "stock", "FRW billion"),
    "Foreign currency (million USD)": ("Central government external debt", "stock", "USD million"),
    "Foreign currency (Billion FRW)": ("Central government external debt", "stock", "FRW billion"),
    "SOE'S loans(Billion FRW)": ("SOE loans", "stock", "FRW billion"),
    "Domestic debt( in Billion Rwf)": ("General government domestic debt", "stock", "FRW billion"),
    "O/W:Medium and long term": ("General government domestic debt: medium and long term", "stock", "FRW billion"),
    "Short term@domestic": ("General government domestic debt: short term", "stock", "FRW billion"),
    "Domestic debt( in Million USD)": ("General government domestic debt", "stock", "USD million"),
}
_CREDITORS = {"Multilaterals", "Bilaterals", "Commercial"}
# Published gaps between the creditor rows and the CG external total (USD m),
# kept as printed; any OTHER gap, or a change to these, raises.
_KNOWN_CREDITOR_GAPS: dict[int, float] = {2021: 10.2}


def _num(v):
    return float(v) if isinstance(v, (int, float)) else None


def parse(path: str, extras: list[str] | None = None) -> pd.DataFrame:
    ws = openpyxl.load_workbook(path, data_only=True).active
    years = {c.column: int(c.value) for c in ws[1]
             if isinstance(c.value, (int, float)) and 2000 < c.value < 2100}
    if len(years) < 5:
        raise ValueError(f"MINECOFIN: year header {years}")
    out, seen, block = [], {}, None
    for r in ws.iter_rows(min_row=2):
        lab = r[0].value
        if not isinstance(lab, str):
            continue
        lab = " ".join(lab.split())
        if lab.startswith("Domestic debt( in Billion"):
            block = "domestic"
        elif lab == "Public External Debt":
            block = "pubext"
        elif lab.startswith("by maturity") or lab.startswith("Domestic debt( in Million"):
            block = None if block == "pubext" else block
        key = lab
        if lab == "Short term" and block == "domestic":
            key = "Short term@domestic"
        if key in _ROWS and key not in seen:
            spec = _ROWS[key]
        elif block == "pubext" and lab in _CREDITORS and lab not in seen:
            spec = (f"Central government external debt: {lab.lower()}", "stock", "USD million")
            key = lab
        else:
            continue
        seen[key] = {}
        comp, meas, unit = spec
        for col, y in years.items():
            v = _num(r[col - 1].value)
            if v is None:
                continue
            seen[key][y] = v
            out.append({"component": comp, "period": str(y), "frequency": "annual",
                        "measure": meas, "value": round(v, 9), "unit": unit,
                        "series_code": "MINECOFIN Debt Stock workbook"})
        if key == "Short term@domestic":
            block = None
    missing = [k for k in _ROWS if k not in seen] + [c for c in _CREDITORS if c not in seen]
    if missing:
        raise ValueError(f"MINECOFIN: rows not found {missing}")
    for y in years.values():
        tot = seen["Foreign currency (million USD)"].get(y)
        parts = [seen[c].get(y) for c in _CREDITORS]
        if tot is not None and None not in parts:
            gap = round(sum(parts) - tot, 1)
            if abs(gap) > 0.5 and _KNOWN_CREDITOR_GAPS.get(y) != gap:
                raise ValueError(f"MINECOFIN {y}: creditors {sum(parts)} != {tot} "
                                 f"(gap {gap})")
        d, a, b = (seen[k].get(y) for k in ("Domestic debt( in Billion Rwf)",
                                            "O/W:Medium and long term", "Short term@domestic"))
        if None not in (d, a, b) and abs(a + b - d) > 0.5:
            raise ValueError(f"MINECOFIN {y}: domestic {a}+{b} != {d}")
    return pd.DataFrame(out)
